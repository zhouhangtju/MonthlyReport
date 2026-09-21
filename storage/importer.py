"""Import exported CSV/XLSX data with archiving, deduplication and audit logs."""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
import tempfile
import re
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from config.datasets import Dataset, get_dataset
from storage.database import connect, initialize
from storage.raw_tables import quote_identifier, database_columns, ensure_columns, validate_headers, version_table


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def normalize_cell(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return value


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def row_payload(row: dict[str, object]) -> tuple[str, str]:
    payload = json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return payload, hashlib.sha256(payload.encode("utf-8")).hexdigest()


def iter_csv(path: Path) -> Iterator[dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if not reader.fieldnames:
            raise ValueError("CSV 缺少表头")
        for row in reader:
            yield {str(key).strip(): normalize_cell(value) for key, value in row.items()}


def iter_excel(path: Path) -> Iterator[dict[str, object]]:
    with path.open("rb") as probe:
        signature = probe.read(8)
    if signature.startswith(b"\xd0\xcf\x11\xe0"):
        yield from iter_legacy_xls(path)
        return
    # 使用文件对象可兼容扩展名为 .xls、实际内容为 XLSX 的平台导出文件。
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                raise ValueError("Excel 缺少表头")
            columns = [str(value).strip() if value is not None else "" for value in header]
            if not all(columns) or len(columns) != len(set(columns)):
                raise ValueError("Excel 表头包含空列名或重复列名")
            for values in rows:
                row = {
                    column: normalize_cell(values[index] if index < len(values) else None)
                    for index, column in enumerate(columns)
                }
                if any(value not in (None, "") for value in row.values()):
                    yield row
        finally:
            workbook.close()


def iter_legacy_xls(path: Path) -> Iterator[dict[str, object]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("读取传统 XLS 需要 xlrd，请安装 requirements.txt") from exc
    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        raise ValueError("Excel 缺少表头")
    columns = [str(sheet.cell_value(0, index)).strip() for index in range(sheet.ncols)]
    if not all(columns) or len(columns) != len(set(columns)):
        raise ValueError("Excel 表头包含空列名或重复列名")
    for row_index in range(1, sheet.nrows):
        row: dict[str, object] = {}
        for column_index, column in enumerate(columns):
            cell = sheet.cell(row_index, column_index)
            value: object = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, workbook.datemode)
            row[column] = normalize_cell(value)
        if any(value not in (None, "") for value in row.values()):
            yield row


def iter_rows(path: Path) -> Iterator[dict[str, object]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        yield from iter_excel(path)
    elif suffix in {".csv", ".tsv"}:
        if suffix == ".tsv":
            raise ValueError("TSV 暂未支持，请先转换为 CSV")
        yield from iter_csv(path)
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")


def archive_file(path: Path, archive_root: Path, dataset: Dataset, digest: str) -> Path:
    day = datetime.now().strftime("%Y-%m-%d")
    target_dir = archive_root / dataset.code / day
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{digest[:12]}_{path.name}"
    if not target.exists():
        shutil.copy2(path, target)
    return target.resolve()


def read_headers(path: Path) -> list[str]:
    """Read even an empty worksheet's header; reject ambiguous layouts."""
    if path.suffix.lower() == '.csv':
        with path.open('r', encoding='utf-8-sig', newline='') as stream:
            values = next(csv.reader(stream), [])
    else:
        with path.open('rb') as stream:
            signature = stream.read(8)
        if signature.startswith(b'\xd0\xcf\x11\xe0'):
            import xlrd
            workbook = xlrd.open_workbook(path)
            try:
                sheet = workbook.sheet_by_index(0)
                values = sheet.row_values(0) if sheet.nrows else []
            finally:
                workbook.release_resources()
        else:
            with path.open('rb') as stream:
                workbook = load_workbook(stream, read_only=True, data_only=True)
                try:
                    sheet = workbook.worksheets[0]
                    sheet.reset_dimensions()
                    values = next(sheet.iter_rows(values_only=True), [])
                finally:
                    workbook.close()
    headers = [str(v).strip() if v is not None else '' for v in values]
    if not headers or not all(headers) or len({h.casefold() for h in headers}) != len(headers):
        raise ValueError('文件表头为空，或包含空列名、重复列名')
    return headers


def _db_value(value: object) -> str | None:
    value = normalize_cell(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return '1' if value else '0'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _stage(source_path, dataset, staged, counts, errors, columns):
    validate_headers(dataset.code, columns)
    seen = {}
    content = hashlib.sha256()
    for number, raw in enumerate(iter_rows(source_path), 2):
        counts['read'] += 1
        row = {column: _db_value(raw.get(column)) for column in columns}
        key = None
        if not dataset.auto_increment:
            key = row[dataset.key_column]
            if key is None or not key.strip() or len(key) > 255:
                counts['failed'] += 1
                errors.append((number, None, f'主键 {dataset.key_column} 为空或超过255字符'))
                continue
            # Primary identifiers use the same whitespace normalization as before.
            row[dataset.key_column] = key = key.strip()
        payload, row_hash = row_payload(row)
        if not dataset.auto_increment and key in seen:
            prior_hash, prior_line = seen[key]
            if prior_hash != row_hash:
                counts['failed'] += 1
                errors.append((number, key, f'与第{prior_line}行主键相同但内容不同，拒绝覆盖'))
            else:
                counts['duplicates'] += 1
            continue
        if not dataset.auto_increment:
            seen[key] = (row_hash, number)
        staged.write(json.dumps([number, row, row_hash], ensure_ascii=False) + '\n')
        content.update(payload.encode('utf-8') + b'\n')
    if errors:
        raise ValueError(f'{dataset.code}: {len(errors)} 行主键为空或冲突，整批未写入；详见 raw_import_error')
    staged.seek(0)
    return content.hexdigest()


def _audit_and_version(connection, dataset, key, row, row_hash, run_id, observed, existing=None):
    payload, _ = row_payload(row)
    first_run = existing['first_run_id'] if existing else run_id
    first_seen = existing['first_seen_at'] if existing else observed
    connection.execute(
        '''INSERT INTO raw_record_audit
           (dataset_code, source_record_id, row_hash, first_run_id, last_run_id, first_seen_at, last_seen_at)
           VALUES (?, ?, ?, ?, ?, ?, ?)
           ON DUPLICATE KEY UPDATE row_hash=VALUES(row_hash), last_run_id=VALUES(last_run_id), last_seen_at=VALUES(last_seen_at)''',
        (dataset.code, key, row_hash, first_run, run_id, first_seen, observed),
    )
    connection.execute(
        f'''INSERT INTO {quote_identifier(version_table(dataset.code))}
            (source_record_id, run_id, row_hash, source_data, observed_at) VALUES (?, ?, ?, ?, ?)
            ON DUPLICATE KEY UPDATE source_record_id=VALUES(source_record_id)''',
        (key, run_id, row_hash, payload, observed),
    )


def _import_keyed(connection, dataset, staged, run_id, counts):
    columns = database_columns(connection, dataset.code)
    table = quote_identifier(dataset.code)
    names = ', '.join(map(quote_identifier, columns))
    insert = f"INSERT INTO {table} ({names}) VALUES ({', '.join('?' for _ in columns)})"
    for line in staged:
        _number, row, row_hash = json.loads(line)
        key = row[dataset.key_column]
        existing_row = connection.execute(
            f'SELECT * FROM {table} WHERE {quote_identifier(dataset.key_column)}=?', (key,)
        ).fetchone()
        audit = connection.execute(
            'SELECT * FROM raw_record_audit WHERE dataset_code=? AND source_record_id=?',
            (dataset.code, key),
        ).fetchone()
        row = {c: row[c] if c in row else (_db_value(existing_row[c]) if existing_row else None) for c in columns}
        _, row_hash = row_payload(row)
        if existing_row is None:
            connection.execute(insert, tuple(row[c] for c in columns))
            counts['inserted'] += 1
        elif row_payload({c: _db_value(existing_row[c]) for c in columns})[1] == row_hash:
            counts['unchanged'] += 1
        else:
            assignments = ', '.join(f'{quote_identifier(c)}=?' for c in columns)
            connection.execute(f'UPDATE {table} SET {assignments} WHERE {quote_identifier(dataset.key_column)}=?',
                               (*[row[c] for c in columns], key))
            counts['updated'] += 1
        _audit_and_version(connection, dataset, key, row, row_hash, run_id, utc_now(), audit)


def _import_period(connection, dataset, staged, content_hash, start, end, run_id, counts):
    period = (dataset.code, start, end)
    current = connection.execute(
        'SELECT * FROM raw_period_batch WHERE dataset_code=? AND period_start=? AND period_end=? FOR UPDATE', period
    ).fetchone()
    table = quote_identifier(dataset.code)
    columns = database_columns(connection, dataset.code)
    # Compare actual business cells, including NULL versus empty string, rather
    # than trusting historical hashes or MySQL's case-insensitive collation.
    existing_rows = {}
    cursor = connection.execute(f'SELECT * FROM {table} ORDER BY id')
    for existing in cursor:
        values = tuple(existing[c] for c in columns)
        existing_rows.setdefault(values, existing['id'])
    if current:
        counts['replaced'] = current['row_count']
        connection.execute('DELETE FROM raw_period_row WHERE dataset_code=? AND period_start=? AND period_end=?', period)
    connection.execute(
        """INSERT INTO raw_period_batch (dataset_code, period_start, period_end, run_id, content_hash, row_count)
           VALUES (?, ?, ?, ?, ?, 0) ON DUPLICATE KEY UPDATE run_id=VALUES(run_id),
           content_hash=VALUES(content_hash), row_count=0""",
        (*period, run_id, content_hash),
    )
    names = ', '.join(map(quote_identifier, columns))
    insert = f"INSERT INTO {table} ({names}) VALUES ({', '.join('?' for _ in columns)})"
    linked = set()
    for line in staged:
        _number, row, row_hash = json.loads(line)
        row = {c: row.get(c) for c in columns}
        _, row_hash = row_payload(row)
        values = tuple(row[c] for c in columns)
        row_id = existing_rows.get(values)
        if row_id is None:
            row_id = connection.execute(insert, values).lastrowid
            existing_rows[values] = row_id
            counts['inserted'] += 1
        elif row_id in linked:
            counts['duplicates'] += 1
            continue
        else:
            counts['unchanged'] += 1
        linked.add(row_id)
        connection.execute(
            'INSERT INTO raw_period_row (dataset_code, period_start, period_end, source_row_number, row_id) VALUES (?, ?, ?, ?, ?)',
            (*period, len(linked), row_id),
        )
        audit = connection.execute(
            'SELECT * FROM raw_record_audit WHERE dataset_code=? AND source_record_id=?',
            (dataset.code, str(row_id)),
        ).fetchone()
        _audit_and_version(connection, dataset, str(row_id), row, row_hash, run_id, utc_now(), audit)
    connection.execute(
        'UPDATE raw_period_batch SET row_count=? WHERE dataset_code=? AND period_start=? AND period_end=?',
        (len(linked), *period),
    )


def import_file(database_path: Path | None, dataset: Dataset, source_path: Path,
                archive_root: Path | None = None, period_start: str | None = None,
                period_end: str | None = None) -> dict[str, object]:
    dataset = get_dataset(dataset.code)
    source_path = source_path.expanduser().resolve()
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    if dataset.auto_increment and (not period_start or not period_end):
        raise ValueError('出入库文件必须提供 period_start 和 period_end，按周期替换而非重复追加')
    if (period_start is None) != (period_end is None):
        raise ValueError('开始日期和结束日期必须一起提供')
    if period_start and date.fromisoformat(period_start) > date.fromisoformat(period_end):
        raise ValueError('开始日期不能晚于结束日期')
    initialize(database_path)
    digest = file_sha256(source_path)
    archived = archive_file(source_path, archive_root, dataset, digest) if archive_root else None
    run_id = f"etl_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}"
    counts = dict(read=0, inserted=0, updated=0, unchanged=0, failed=0, duplicates=0, replaced=0)
    errors = []
    with connect(database_path) as connection:
        connection.execute(
            '''INSERT INTO etl_run (run_id, dataset_code, source_system, source_file, archived_file,
               file_sha256, period_start, period_end, started_at, status)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'running')''',
            (run_id, dataset.code, dataset.source_system, str(source_path), str(archived) if archived else None,
             digest, period_start, period_end, utc_now()),
        )
        connection.commit()
        database_name = connection.execute('SELECT DATABASE() AS name').fetchone()['name']
        lock_name = 'raw:' + hashlib.sha256(f'{database_name}:{dataset.code}'.encode()).hexdigest()[:50]
        locked = False
        try:
            locked = connection.execute('SELECT GET_LOCK(?, 60) AS acquired', (lock_name,)).fetchone()['acquired'] == 1
            if not locked:
                raise RuntimeError(f'{dataset.code}: 等待其他导入结束超时')
            with tempfile.TemporaryFile(mode='w+', encoding='utf-8') as staged:
                headers = read_headers(source_path)
                content_hash = _stage(source_path, dataset, staged, counts, errors, headers)
                if file_sha256(source_path) != digest:
                    raise RuntimeError('校验期间源文件发生变化，请重新导入')
                ensure_columns(connection, dataset.code, headers)
                if dataset.auto_increment:
                    _import_period(connection, dataset, staged, content_hash, period_start, period_end, run_id, counts)
                else:
                    _import_keyed(connection, dataset, staged, run_id, counts)
            connection.execute(
                '''UPDATE etl_run SET finished_at=?, status='success', rows_read=?, rows_inserted=?,
                   rows_updated=?, rows_unchanged=?, rows_failed=?, rows_duplicates=? WHERE run_id=?''',
                (utc_now(), counts['read'], counts['inserted'], counts['updated'], counts['unchanged'],
                 counts['failed'], counts['duplicates'], run_id),
            )
            connection.commit()
        except Exception as exc:
            connection.rollback()
            connection.execute(
                '''UPDATE etl_run SET finished_at=?, status='failed', error_message=?, rows_read=?,
                   rows_failed=?, rows_duplicates=? WHERE run_id=?''',
                (utc_now(), str(exc), counts['read'], max(counts['failed'], 1), counts['duplicates'], run_id),
            )
            if errors:
                connection.executemany(
                    'INSERT INTO raw_import_error (run_id, source_row_number, source_record_id, error_message) VALUES (?, ?, ?, ?)',
                    [(run_id, *item) for item in errors],
                )
            connection.commit()
            raise
        finally:
            if locked:
                connection.execute('SELECT RELEASE_LOCK(?)', (lock_name,))
    return {'run_id': run_id, 'archived_file': str(archived) if archived else None, **counts}
