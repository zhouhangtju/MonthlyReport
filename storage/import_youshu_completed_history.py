"""从历史工作簿 Sheet1 导入地市月度装机量；默认只校验，--apply 才写 MySQL。"""
import argparse
import calendar
from collections import defaultdict
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import uuid

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from storage.database import connect
from storage.youshu_install_completed import CITIES, count_value, now

DATASET = 'youshu_install_completed_history_city'


def parse(path, sheet='Sheet1'):
    path = Path(path).resolve()
    content = path.read_bytes()
    workbook = load_workbook(BytesIO(content), data_only=False)
    groups = defaultdict(dict)
    totals = {}
    try:
        ws = workbook[sheet]
        for offset in (0, 4):
            if [ws.cell(2, offset+i).value for i in (1,2,3)] != ['月份','地市','工单量']:
                raise ValueError('Sheet1 两个数据块表头应为 月份、地市、工单量')
            for line in range(3, ws.max_row+1):
                cells = [ws.cell(line, offset+i) for i in (1,2,3)]
                raw, city, value = [c.value for c in cells]
                if all(v is None for v in (raw,city,value)):
                    continue
                match = re.fullmatch(r'(\d{2}|\d{4})年(\d{1,2})月', str(raw).strip())
                if not match:
                    raise ValueError(f'{cells[0].coordinate} 月份格式错误')
                year, month = map(int, match.groups())
                if year < 100:
                    year += 2000
                last = calendar.monthrange(year, month)[1]
                bounds = (f'{year:04}-{month:02}-01', f'{year:04}-{month:02}-{last:02}')
                city = str(city).strip()
                if city == '全省':
                    if bounds in totals:
                        raise ValueError(f'{bounds} 全省合计重复')
                    totals[bounds] = (cells[2].coordinate, value)
                else:
                    if city not in CITIES or city in groups[bounds]:
                        raise ValueError(f'{bounds} 地市未知或重复：{city}')
                    groups[bounds][city] = {'value': count_value(value), 'cell': cells[2].coordinate}
        if not groups or set(groups) != set(totals):
            raise ValueError('月份数据或全省合计缺失')
        for bounds, rows in groups.items():
            if set(rows) != CITIES:
                raise ValueError(f'{bounds} 缺少地市')
            coordinate, total = totals[bounds]
            if isinstance(total,str) and total.startswith('='):
                # Only the exact local SUM of the month's 11 raw cells is permitted.
                match = re.fullmatch(r'=SUM\(([A-Z]+)(\d+):\1(\d+)\)',total,re.I)
                if not match:
                    raise ValueError(f'{coordinate} 不支持的合计公式')
                col, first, last = match.groups()
                refs = {f'{col.upper()}{i}' for i in range(int(first),int(last)+1)}
                if refs != {r['cell'] for r in rows.values()}:
                    raise ValueError(f'{coordinate} SUM范围不是本月11地市')
                total = sum(r['value'] for r in rows.values())
            if count_value(total) != sum(r['value'] for r in rows.values()):
                raise ValueError(f'{bounds} 地市与全省合计不一致')
    finally:
        workbook.close()
    return {'path': str(path), 'sheet': sheet, 'sha256': hashlib.sha256(content).hexdigest(),
            'months': dict(sorted(groups.items()))}


def import_history(parsed, apply=False, overwrite=False):
    """Overwrite replaces active snapshots atomically, retaining old rows for audit."""
    plans = []
    with connect() as conn:
        for (start,end), rows in parsed['months'].items():
            existing = conn.execute('''SELECT c.collection_run_id FROM collection_run c
                WHERE c.status='success' AND c.period_start=? AND c.period_end=?
                AND c.dataset_code IN ('youshu_install_completed',?)
                AND EXISTS (SELECT 1 FROM youshu_install_completed_city d
                            WHERE d.collection_run_id=c.collection_run_id)
                ORDER BY c.finished_at DESC,c.collection_run_id DESC LIMIT 1 FOR UPDATE''',
                (start,end,DATASET)).fetchone()
            action = 'insert'
            if existing:
                values = conn.execute('''SELECT city,completed_order_count FROM youshu_install_completed_city
                    WHERE collection_run_id=?''',(existing['collection_run_id'],)).fetchall()
                actual = {r['city']: int(r['completed_order_count']) for r in values}
                if len(values) != 11 or actual != {c:r['value'] for c,r in rows.items()}:
                    if not overwrite:
                        raise ValueError(f'{start} 至 {end} 与数据库已有批次 {existing["collection_run_id"]} 不一致，整批未写入；确认替换时使用 --apply --overwrite')
                    action = 'overwrite'
                else:
                    action = 'skip_identical'
            plans.append({'start': start,'end': end,'action': action,
                          'previous_run_id': existing['collection_run_id'] if existing else None,
                          'total':sum(r['value'] for r in rows.values())})
        if apply:
            for plan in plans:
                if plan['action'] == 'skip_identical':
                    continue
                start,end = plan['start'],plan['end']
                rows = parsed['months'][(start,end)]
                run_id = f'history_city_{start}_{uuid.uuid4().hex[:24]}'
                etl_id = run_id+'_etl'
                stamp = now()
                provenance = json.dumps({'kind':'historical_manual_city_only', 'file':parsed['path'],
                    'sheet':parsed['sheet'],'sha256':parsed['sha256'],
                    'operation':plan['action'], 'previous_run_id':plan['previous_run_id'],
                    'cells':{c:r['cell'] for c,r in rows.items()}},ensure_ascii=False)
                if plan['action'] == 'overwrite':
                    # Invalidate all old successful same-period snapshots, including paired online runs.
                    # Their raw city/county rows and ETL audit remain intact for historical traceability.
                    conn.execute('''UPDATE collection_run SET status='superseded'
                        WHERE status='success' AND period_start=? AND period_end=?
                        AND dataset_code IN ('youshu_install_completed',?)''',(start,end,DATASET))
                conn.execute('''INSERT INTO collection_run (collection_run_id,dataset_code,period_start,period_end,
                    started_at,finished_at,status,expected_files,produced_files,imported_files,source_files,etl_run_ids)
                    VALUES (?,?,?,?,?,?,'success',1,1,1,?,?)''',
                    (run_id,DATASET,start,end,stamp,stamp,provenance,json.dumps([etl_id])))
                conn.execute('''INSERT INTO etl_run (run_id,dataset_code,source_system,source_file,file_sha256,
                    period_start,period_end,started_at,finished_at,status,rows_read,rows_inserted)
                    VALUES (?,?,?,?,?,?,?,?,?,'success',12,11)''',
                    (etl_id,'youshu_install_completed_city','历史人工补录',parsed['path'],parsed['sha256'],start,end,stamp,stamp))
                conn.executemany('''INSERT INTO youshu_install_completed_city (collection_run_id,period_start,
                    period_end,city,completed_order_count,source_file,collected_at) VALUES (?,?,?,?,?,?,?)''',
                    [(run_id,start,end,c,r['value'],parsed['path'],stamp) for c,r in rows.items()])
                plan['collection_run_id'] = run_id
    return {'applied':apply,'overwrite':overwrite,
            'inserted_rows':sum(p['action'] in ('insert','overwrite') for p in plans)*11 if apply else 0,'plan':plans}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--file',type=Path,required=True)
    parser.add_argument('--sheet',default='Sheet1')
    parser.add_argument('--overwrite',action='store_true',help='用Excel替换冲突月份的有效地市批次，旧批次保留为superseded')
    mode=parser.add_mutually_exclusive_group()
    mode.add_argument('--check-db',action='store_true',help='只检查数据库冲突，不写入')
    mode.add_argument('--apply',action='store_true',help='检查冲突并事务写入已有MySQL表')
    args=parser.parse_args()
    if args.overwrite and not (args.apply or args.check_db):
        parser.error('--overwrite 必须配合 --apply 或 --check-db')
    parsed=parse(args.file,args.sheet)
    if args.apply or args.check_db:
        result=import_history(parsed,args.apply,args.overwrite)
    else:
        result={'validated':True,'months':len(parsed['months']),'city_rows':sum(len(r) for r in parsed['months'].values()),
                'periods':[{'start':s,'end':e,'total':sum(r['value'] for r in rows.values())} for (s,e),rows in parsed['months'].items()]}
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
