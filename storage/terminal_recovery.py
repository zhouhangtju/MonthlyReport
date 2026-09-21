"""Lossless terminal recovery source snapshots and collection coverage."""

import hashlib
import json
import uuid
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from storage.database import connect, initialize as initialize_database


SOURCE_FILES = {
    "integration_removal_order": "一体化专线拆机清单_{month}.xlsx",
    "eoms_service_removal_order": "服务类产品支撑工单_{month}.xlsx",
    "integration_terminal_inbound": "终端入库_{month}.xlsx",
    "integration_material_baseline": "全省物资基准库.xlsx",
    "terminal_material_names": "物料名称表.xlsx",
}

REPORT_SHEETS = (
    "物料名称表", "全省物资基准库", "一体化专线拆机清单（不含小微版）",
    "EOMS服务类拆机清单", "一体化拆回设备清单", "拆除清单汇总",
    "按地市回收率汇总", "匹配后拆回设备清单（三类标签已去重）",
    "sheet2", "sheet3", "sheet1",
)


def encode_cell(value):
    if isinstance(value, datetime):
        return {"datetime": value.isoformat()}
    if isinstance(value, date):
        return {"date": value.isoformat()}
    return value


def decode_cell(value):
    if isinstance(value, dict):
        if "datetime" in value:
            return datetime.fromisoformat(value["datetime"])
        if "date" in value:
            return date.fromisoformat(value["date"])
    return value


def save_source(database, code, source, start, end, etl_run_id):
    """Store full ordered cell data plus the original workbook for formatting."""
    source = Path(source)
    content = source.read_bytes()
    snapshot_id = uuid.uuid4().hex
    wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        ws = wb.active
        initialize_database(database)
        with connect(database) as conn:
            conn.execute(
                """INSERT INTO tr_source_snapshot (
                       snapshot_id, dataset_code, period_start, period_end,
                       etl_run_id, created_at, filename, sha256, row_count, workbook
                   ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    snapshot_id, code, start, end, etl_run_id,
                    datetime.now().isoformat(), source.name,
                    hashlib.sha256(content).hexdigest(), ws.max_row, content,
                ),
            )
            conn.executemany(
                """INSERT INTO tr_source_row (
                       snapshot_id, source_row_number, cells_json
                   ) VALUES (?, ?, ?)""",
                (
                    (
                        snapshot_id,
                        number,
                        json.dumps([encode_cell(v) for v in row], ensure_ascii=False),
                    )
                    for number, row in enumerate(ws.iter_rows(values_only=True), 1)
                ),
            )
    finally:
        wb.close()
    return snapshot_id


def restore_sources(database, start, end, output_dir):
    """Select exact-period snapshots, never fall back to raw latest-state records."""
    initialize_database(database)
    output_dir = Path(output_dir)
    snapshots = {}
    with connect(database) as conn:
        for code in SOURCE_FILES:
            row = conn.execute("""SELECT * FROM tr_source_snapshot
                WHERE dataset_code=? AND period_start=? AND period_end=?
                ORDER BY created_at DESC, snapshot_id DESC LIMIT 1""", (code, start, end)).fetchone()
            if row is None:
                raise ValueError(f"Missing source snapshot: {code} ({start}..{end})")
            snapshots[code] = row
        output_dir.mkdir(parents=True, exist_ok=True)
        for code, snapshot in snapshots.items():
            workbook = bytes(snapshot["workbook"])
            if hashlib.sha256(workbook).hexdigest() != snapshot["sha256"]:
                raise ValueError(f"Corrupted source snapshot: {code}")
            wb = load_workbook(BytesIO(workbook), data_only=False)
            try:
                ws = wb.active
                count = 0
                rows = conn.execute(
                    """SELECT source_row_number, cells_json
                       FROM tr_source_row
                       WHERE snapshot_id=?
                       ORDER BY source_row_number""",
                    (snapshot["snapshot_id"],),
                )
                for record in rows:
                    number = record["source_row_number"]
                    cells_json = record["cells_json"]
                    count += 1
                    if number != count:
                        raise ValueError(f"Non-contiguous source rows: {code}")
                    for column, value in enumerate(json.loads(cells_json), 1):
                        cell = ws.cell(number, column)
                        # Merged-cell placeholders have no writable value.
                        if cell.__class__.__name__ != "MergedCell":
                            cell.value = decode_cell(value)
                if count != snapshot["row_count"]:
                    raise ValueError(f"Incomplete source snapshot: {code}")
                wb.save(output_dir / SOURCE_FILES[code].format(month=start[:7]))
            finally:
                wb.close()
    return {code: row["snapshot_id"] for code, row in snapshots.items()}


def source_run_ids(database, snapshots):
    initialize_database(database)
    with connect(database) as conn:
        return list(dict.fromkeys(
            conn.execute(
                "SELECT etl_run_id FROM tr_source_snapshot WHERE snapshot_id=?",
                (snapshot,),
            ).fetchone()["etl_run_id"]
            for snapshot in snapshots.values()
        ))


def has_source_coverage(database, codes, start, end):
    try:
        initialize_database(database)
        with connect(database) as conn:
            return all(conn.execute(
                "SELECT 1 FROM tr_source_snapshot WHERE dataset_code=? AND period_start=? AND period_end=? LIMIT 1",
                (code, start, end),
            ).fetchone() for code in codes)
    except Exception:
        return False
