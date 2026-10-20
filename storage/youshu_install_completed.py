"""Exact-period installation completion snapshots; independent of raw order tables."""
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from storage.database import connect

DATASET = "youshu_install_completed"
TABLES = {level: f"{DATASET}_{level}" for level in ("city", "county")}
CITIES = set("杭州 湖州 嘉兴 金华 丽水 宁波 绍兴 台州 温州 舟山 衢州".split())


def initialize():
    """Add only the new tables; do not run unrelated legacy migrations."""
    sql = Path(__file__).resolve().parents[1] / "sql/create_youshu_install_completed.sql"
    with connect() as conn:
        # These shared audit tables are supplied by the existing application schema.
        conn.execute("SELECT collection_run_id FROM collection_run LIMIT 0")
        conn.execute("SELECT run_id FROM etl_run LIMIT 0")
        conn.executescript(sql.read_text("utf-8"))


def now():
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


def validate_period(start, end):
    if date.fromisoformat(start).isoformat() != start or date.fromisoformat(end).isoformat() != end:
        raise ValueError("日期必须为 YYYY-MM-DD")
    if start > end:
        raise ValueError("开始日期不能晚于结束日期")


def count_value(value):
    if value is None or isinstance(value, bool):
        raise ValueError("工单量不能为空或布尔值")
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"工单量不是数字：{value}") from exc
    if not number.is_finite() or number < 0 or number != number.to_integral_value() or number > 18446744073709551615:
        raise ValueError(f"工单量必须为非负 BIGINT 整数：{value}")
    return int(number)


def parse_file(path, level):
    if level not in TABLES:
        raise ValueError("level 必须是 city 或 county")
    path = Path(path)
    expected = ["地市", "装机报结工单量"] if level == "city" else ["地市", "区县", "装机报结工单量"]
    workbook = load_workbook(path, read_only=True, data_only=False)
    rows, keys, total = [], set(), None
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("报结量导出必须只有一个工作表")
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True)
        if list(next(iterator, ())) != expected:
            raise ValueError(f"{level} 表头必须是 {expected}")
        for line, values in enumerate(iterator, 2):
            if all(v is None for v in values):
                continue
            if len(values) != len(expected):
                raise ValueError(f"第 {line} 行列数不符")
            city = str(values[0]).strip() if values[0] is not None else ""
            amount = count_value(values[-1])
            county = str(values[1]).strip() if level == "county" and values[1] is not None else ""
            if city == "总计":
                if total is not None or (level == "county" and county):
                    raise ValueError("总计行重复或结构异常")
                total = amount
                continue
            if city not in CITIES or (level == "county" and not county):
                raise ValueError(f"第 {line} 行地区为空或未知：{city}/{county}")
            key = (city, county)
            if key in keys:
                raise ValueError(f"地区重复：{key}")
            keys.add(key)
            rows.append({"city": city, "county": county, "count": amount})
    finally:
        workbook.close()
    if {r["city"] for r in rows} != CITIES:
        raise ValueError("未覆盖全省 11 个地市，不能将缺失地区视为零")
    if total is None or sum(r["count"] for r in rows) != total:
        raise ValueError("缺少平台总计或地区合计与总计不一致")
    return {"rows": rows, "total": total, "source_file": str(path.resolve()),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def validate_pair(city_path, county_path):
    parsed = {"city": parse_file(city_path, "city"), "county": parse_file(county_path, "county")}
    sums = defaultdict(int)
    for row in parsed["county"]["rows"]:
        sums[row["city"]] += row["count"]
    if parsed["city"]["total"] != parsed["county"]["total"] or any(
        sums[row["city"]] != row["count"] for row in parsed["city"]["rows"]
    ):
        raise ValueError("地市、区县汇总交叉校验失败")
    return parsed


def latest_success(start, end):
    """A larger period never covers a smaller aggregate period."""
    validate_period(start, end)
    with connect() as conn:
        row = conn.execute("""SELECT c.collection_run_id FROM collection_run c
            WHERE c.dataset_code=? AND c.period_start=? AND c.period_end=?
              AND c.status='success' AND c.expected_files=2
              AND c.produced_files=2 AND c.imported_files=2
              AND EXISTS (SELECT 1 FROM youshu_install_completed_city d
                  WHERE d.collection_run_id=c.collection_run_id)
              AND EXISTS (SELECT 1 FROM youshu_install_completed_county d
                  WHERE d.collection_run_id=c.collection_run_id)
            ORDER BY c.finished_at DESC, c.collection_run_id DESC LIMIT 1""",
            (DATASET, start, end)).fetchone()
    return row["collection_run_id"] if row else None


def load_latest(start, end):
    """Return both levels from the same successful batch, never sum snapshots."""
    run_id = latest_success(start, end)
    if run_id is None:
        raise ValueError("该统计周期没有完整成功的报结量数据")
    result = {"collection_run_id": run_id}
    with connect() as conn:
        for level, table in TABLES.items():
            order = "city,county" if level == "county" else "city"
            result[level] = [dict(row) for row in conn.execute(
                f"SELECT * FROM {table} WHERE collection_run_id=? ORDER BY {order}",
                (run_id,)).fetchall()]
    return result


def load_latest_city(start, end):
    """Calculation-only reader: accept city history without pretending county coverage."""
    validate_period(start, end)
    with connect() as conn:
        batch = conn.execute('''SELECT c.collection_run_id,c.dataset_code FROM collection_run c
            WHERE c.status='success' AND c.period_start=? AND c.period_end=?
              AND ((c.dataset_code='youshu_install_completed' AND c.expected_files=2
                    AND c.produced_files=2 AND c.imported_files=2)
                OR (c.dataset_code='youshu_install_completed_history_city' AND c.expected_files=1
                    AND c.produced_files=1 AND c.imported_files=1))
            ORDER BY c.finished_at DESC,c.collection_run_id DESC LIMIT 1''',(start,end)).fetchone()
        if batch is None:
            raise ValueError('该统计周期没有完整成功的报结量数据')
        rows = conn.execute('''SELECT * FROM youshu_install_completed_city
            WHERE collection_run_id=? ORDER BY city''',(batch['collection_run_id'],)).fetchall()
    return {'collection_run_id': batch['collection_run_id'], 'dataset_code': batch['dataset_code'],
            'city': [dict(r) for r in rows]}


def start_run(run_id, start, end):
    validate_period(start, end)
    with connect() as conn:
        conn.execute("""INSERT INTO collection_run
            (collection_run_id,dataset_code,period_start,period_end,started_at,status,
             expected_files,source_files,etl_run_ids)
            VALUES (?,?,?,?,?,'running',2,'[]','[]')""", (run_id, DATASET, start, end, now()))


def fail_run(run_id, files, error):
    with connect() as conn:
        conn.execute("""UPDATE collection_run SET status='failed',finished_at=?,
            produced_files=?,source_files=?,error_message=?
            WHERE collection_run_id=? AND status='running'""",
            (now(), len(files), json.dumps(files, ensure_ascii=False), str(error), run_id))


def import_pair(run_id, start, end, parsed):
    """Both business tables, ETL records and success marker commit together."""
    validate_period(start, end)
    timestamp = now()
    with connect() as conn:
        run = conn.execute("SELECT * FROM collection_run WHERE collection_run_id=? FOR UPDATE", (run_id,)).fetchone()
        if not run or run["dataset_code"] != DATASET or run["period_start"] != start or run["period_end"] != end:
            raise ValueError("采集批次与数据周期不匹配")
        if run["status"] == "success":
            etl_ids = json.loads(run["etl_run_ids"])
            for level, etl_id in zip(TABLES, etl_ids):
                previous = conn.execute("SELECT file_sha256 FROM etl_run WHERE run_id=?", (etl_id,)).fetchone()
                if not previous or previous["file_sha256"] != parsed[level]["sha256"]:
                    raise ValueError("成功批次不可替换为其他源文件，请创建新的采集批次")
            if len(etl_ids) != 2:
                raise ValueError("成功批次的 ETL 审计记录不完整")
            return etl_ids
        if run["status"] != "running":
            raise ValueError("只能向运行中的批次写入数据")
        etl_ids = []
        for level, table in TABLES.items():
            data = parsed[level]
            etl_id = f"{run_id}_{level}"
            etl_ids.append(etl_id)
            conn.execute("""INSERT INTO etl_run
                (run_id,dataset_code,source_system,source_file,file_sha256,period_start,
                 period_end,started_at,finished_at,status,rows_read,rows_inserted)
                VALUES (?,?,?,?,?,?,?,?,?,'success',?,?)""",
                (etl_id, table, "有数", data["source_file"], data["sha256"], start, end,
                 timestamp, timestamp, len(data["rows"])+1, len(data["rows"])))
            columns = "collection_run_id,period_start,period_end,city,completed_order_count,source_file,collected_at"
            values = [(run_id, start, end, r["city"], r["count"], data["source_file"], timestamp) for r in data["rows"]]
            if level == "county":
                columns += ",county"
                values = [v + (r["county"],) for v, r in zip(values, data["rows"])]
            conn.executemany(f"INSERT INTO {table} ({columns}) VALUES ({','.join('?' for _ in values[0])})", values)
        conn.execute("""UPDATE collection_run SET status='success',finished_at=?,
            produced_files=2,imported_files=2,source_files=?,etl_run_ids=?
            WHERE collection_run_id=?""", (timestamp,
            json.dumps([parsed[k]["source_file"] for k in TABLES], ensure_ascii=False), json.dumps(etl_ids), run_id))
    return etl_ids
