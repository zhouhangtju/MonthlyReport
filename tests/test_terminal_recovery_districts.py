import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook

from metrics.terminal_recovery.districts import (
    DistrictMap, clean_row, build_district_report, metric_rows, resolve_mapping,
    NON_EQI_VIDEO_COLUMN, VIDEO_ORDER_COLUMN, EQI_DEVICE_COLUMN,
    EXPECTED_COLUMN, RECOVERED_COLUMN, RATE_COLUMN,
)
from storage.metric_results import FIELDS
from reporting.database_results import terminal_district_chart_rows


class DistrictTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            {"地市": "杭州", "区县": "西湖", "标准区县": "西湖区"},
            {"地市": "杭州", "区县": "余杭", "标准区县": "余杭区"},
            {"地市": "杭州", "区县": "开发", "标准区县": "开发区"},
            {"地市": "宁波", "区县": "开发", "标准区县": "开发区"},
        ]
        self.mapping = DistrictMap(self.rows)

    def test_keyword_ambiguity_and_city(self):
        self.assertEqual(self.mapping.match("开发区客户", "杭州", True), "开发区")
        self.assertEqual(self.mapping.match("开发区客户", "", True), "")
        self.assertEqual(self.mapping.match("西湖区余杭区客户", "杭州", True), "")
        self.assertEqual(self.mapping.match("西湖区客户", "宁波", True), "")

    def test_default_mapping_is_next_to_district_module(self):
        path = resolve_mapping("unused-output-directory")
        self.assertEqual(path.name, "区县信息汇总.xlsx")
        self.assertEqual(path.parent.name, "terminal_recovery")
        self.assertTrue(path.is_file())

    def test_blank_only_business_exclusion(self):
        row = {"地市": "杭州", "业务类型": "行业视频-小微版", "工单主题": "西湖客户"}
        self.assertFalse(clean_row(row, self.mapping)[3])
        row["区县"] = "西湖"
        self.assertEqual(clean_row(row, self.mapping)[1:], ("西湖区", "区县直接映射", True))

    def test_city_value_and_unmapped_clear(self):
        row = {"地市": "杭州", "区县": "杭州市", "工单主题": "西湖区客户"}
        self.assertEqual(clean_row(row, self.mapping)[1], "西湖区")
        row["工单主题"] = "未知客户"
        self.assertEqual(clean_row(row, self.mapping)[1], "")
        row["区县"] = "集团客户中心"
        self.assertFalse(clean_row(row, self.mapping)[3])

    def test_service_priority_and_address_fallback(self):
        row = {"安装地市": "杭州", "所属地市": "杭州市", "所属区县": "西湖",
               "主题": "余杭区客户", "客户机房地址": "西湖区机房"}
        self.assertEqual(clean_row(row, self.mapping, True)[1:3], ("西湖区", "所属区县映射"))
        row["所属地市"] = "宁波"
        self.assertEqual(clean_row(row, self.mapping, True)[1:3], ("余杭区", "主题关键词"))
        row["主题"] = "未知"
        self.assertEqual(clean_row(row, self.mapping, True)[1:3], ("西湖区", "客户机房地址关键词"))

    def test_ppt_bottom_ten_is_displayed_descending_without_city(self):
        rows = [{"metric_value": value, "dimension": {"label": label}}
                for label, value in (("台州·玉环", .82), ("绍兴·越城", .85), ("嘉兴·平湖", .84))]
        displayed, bottom_names = terminal_district_chart_rows(rows)
        self.assertEqual([r["metric_value"] for r in displayed], [.85, .84, .82])
        self.assertEqual([r["dimension"]["label"].partition("·")[2] for r in displayed],
                         ["越城", "平湖", "玉环"])
        self.assertEqual(bottom_names, ["玉环", "平湖", "越城"])

    def test_summary_preserves_rows_quantities_and_zero_denominator(self):
        def write(path, rows, name="Sheet1"):
            wb = Workbook()
            wb.active.title = name
            for row in rows:
                wb.active.append(row)
            wb.save(path)
            wb.close()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            mapping, line, service, recovery, output = [root / f"{name}.xlsx" for name in ("map", "line", "service", "recovery", "result")]
            write(mapping, [["地市", "区县", "标准区县"], *[[r[k] for k in ("地市", "区县", "标准区县")] for r in self.rows]])
            write(line, [["工单号", "地市", "区县", "业务类型"],
                         ["1", "杭州", "西湖", "互联网专线"], ["1", "杭州", "西湖", "互联网专线"],
                         ["2", "杭州", "余杭", "视频监控"],
                         ["4", "杭州", "西湖", "行业视频-小微版"]])
            write(service, [["工单号", "安装地市", "安装区县", "服务类别", "路由器数量", "FTTR数量", "光AP数量"],
                            ["3", "杭州", "开发区", "E企组网", 1, 2, 1],
                            ["5", "杭州", "西湖", "其他服务", 100, 100, 100]])
            write(recovery, [["地市", "区县", "总数量", "型号匹配"], ["杭州", "西湖区", 2, "ONU"], ["杭州", "南城区", 3, "FTTO"], ["宁波", "开发区", 4, "ONU"]], "匹配后拆回设备清单（三类标签已去重）")
            report = build_district_report(line, service, recovery, mapping, output)
            self.assertEqual([r["区县"] for r in report["bottom10"]], ["余杭区", "开发区", "西湖区"])
            west = next(r for r in report["rows"] if r["区县"] == "西湖区")
            self.assertEqual(west[NON_EQI_VIDEO_COLUMN], 2)
            self.assertEqual(west[EXPECTED_COLUMN], 2)
            yuhang = next(r for r in report["rows"] if r["区县"] == "余杭区")
            self.assertEqual(yuhang[VIDEO_ORDER_COLUMN], 1)
            self.assertEqual(yuhang[EXPECTED_COLUMN], 2)
            development = next(r for r in report["rows"] if r["区县"] == "开发区" and r["地市"] == "杭州")
            self.assertEqual(development[EQI_DEVICE_COLUMN], 4)
            self.assertEqual(development[RECOVERED_COLUMN], 3)
            self.assertEqual(development[RATE_COLUMN], .75)
            self.assertIsNone(next(r for r in report["rows"] if r["地市"] == "宁波")[RATE_COLUMN])
            self.assertTrue(metric_rows(report))
            for result in metric_rows(report):
                self.assertTrue(set(result["dimension"]).issubset(FIELDS["terminal_recovery"]))
            wb = load_workbook(output, read_only=True)
            self.assertIn("拆机区县清洗明细", wb.sheetnames)
            wb.close()


if __name__ == "__main__":
    unittest.main()
