import gc
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from collector.eoms import eoms_export_qiwan_auto_login as eoms
from collector.integration import integration_dismantle as materials
from collector.integration import zhuanxian_chaiji_export as lines
from metrics.terminal_recovery.terminal_recovery_export_online import export, validate_dates, unique_rows_by_order


class TerminalRecoveryTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.database = self.root / "test.db"

    def tearDown(self):
        gc.collect()
        self.temp.cleanup()

    def workbook(self, name, header, rows):
        path = self.root / name
        wb = Workbook()
        wb.active.append(header)
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
        return path

    def test_duplicate_details_survive_reimport_and_reordering(self):
        code = "integration_terminal_inbound"
        path = self.workbook("inbound.xlsx", ["物料名称", "总数量"],
                             [["ONU", 1], ["ONU", 1], ["ONU", 2]])
        args = ({code: path}, "2026-08-01", "2026-08-31", "both", self.database)
        for module in (eoms, lines, materials):
            with self.subTest(module=module.__name__):
                database = self.root / f"{module.__name__}.db"
                result = module.finish_exports(*args[:-1], database)["etl"][code]
                self.assertEqual(result["inserted"], 3)
                self.workbook("inbound.xlsx", ["物料名称", "总数量"],
                              [["ONU", 2], ["ONU", 1], ["ONU", 1]])
                result = module.finish_exports(*args[:-1], database)["etl"][code]
                self.assertEqual(result["unchanged"], 3)
                self.assertEqual(result["inserted"], 0)
                self.workbook("inbound.xlsx", ["物料名称", "总数量"],
                              [["ONU", 1], ["ONU", 1], ["ONU", 2]])

    def test_order_collectors_pass_dates_and_import_downloads(self):
        for module, function, code in [
            (lines, "export_order_data", "integration_removal_order"),
            (eoms, "export_qiwan", "eoms_service_removal_order"),
        ]:
            with self.subTest(code=code):
                source = self.workbook(code + ".xlsx", ["工单号"], [["order-1"]])
                with patch.object(module, function, return_value=source) as download, \
                        patch.object(lines, "load_token", return_value=("test-token", "test")):
                    result = module.collect("2026-08-01", "2026-08-31",
                                            database=self.database, output_dir=self.root)
                self.assertEqual(result["etl"][code]["inserted"], 1)
                self.assertTrue(download.call_args.args[0].startswith("2026-08-01"))
                self.assertTrue(download.call_args.args[1].startswith("2026-08-31"))

    def test_online_dedup_keeps_first_filtered_occurrence(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["关联单据号", "物料编码"])
        ws.append(["ZJ-SOC-17-01-260817-1293524", "first"])
        ws.append(["other-order", "other"])
        ws.append(["ZJ-SOC-17-01-260817-1293524", "13652"])
        self.assertEqual(unique_rows_by_order(ws, [2, 3, 4], 1), [2, 3])
        self.assertEqual(unique_rows_by_order(ws, [3, 4], 1), [3, 4])
        wb.close()

    def test_line_token_uses_material_auto_login(self):
        with patch.object(lines.requests, "Session") as session_factory, \
                patch.object(materials, "DEFAULT_ACCOUNT", "test-account"), \
                patch.object(materials, "DEFAULT_PASSWORD", "test-password"), \
                patch.object(materials, "login") as login, \
                patch.object(materials, "extract_token", return_value="fresh-token") as extract, \
                patch.object(materials, "update_cred_json") as save:
            token, account = lines.load_token()
            session = session_factory.return_value.__enter__.return_value
            login.assert_called_once_with(session, "test-account", "test-password")
            extract.assert_called_once_with(login.return_value)
            save.assert_called_once_with("test-account", "fresh-token")
            self.assertEqual((token, account), ("fresh-token", "test-account"))
            session_factory.return_value.__exit__.assert_called_once()

    def test_line_login_failure_stops_download_and_import(self):
        with patch.object(materials, "login", side_effect=RuntimeError("login failed")), \
                patch.object(materials, "update_cred_json") as save, \
                patch.object(lines, "export_order_data") as download, \
                patch.object(lines, "finish_exports") as finish:
            with self.assertRaisesRegex(RuntimeError, "login failed"):
                lines.collect("2026-08-01", "2026-08-31", database=self.database,
                              output_dir=self.root, refresh=True)
            save.assert_not_called()
            download.assert_not_called()
            finish.assert_not_called()

    def test_material_downloads_extend_to_next_month_third(self):
        for start, end, query_end in [
            ("2026-08-01", "2026-08-31", "2026-09-03"),
            ("2026-12-01", "2026-12-31", "2027-01-03"),
            ("2028-02-01", "2028-02-29", "2028-03-03"),
        ]:
            with self.subTest(start=start), ExitStack() as stack:
                stack.enter_context(patch.object(materials, "login"))
                stack.enter_context(patch.object(materials, "extract_token", return_value="test"))
                stack.enter_context(patch.object(materials, "update_cred_json"))
                outbound = stack.enter_context(patch.object(materials, "export_outbound", return_value=True))
                inbound = stack.enter_context(patch.object(materials, "export_inbound", return_value=True))
                stack.enter_context(patch.object(materials, "export_material_baseline", return_value=True))
                finish = stack.enter_context(patch.object(materials, "finish_exports"))
                materials.collect(start, end, database=self.database, output_dir=self.root, refresh=True)
                for download in (outbound, inbound):
                    self.assertEqual(download.call_args.args[2:5], (start[:7], start, query_end))
                self.assertEqual(finish.call_args.args[1:], (start, end, "both", self.database))
                files = finish.call_args.args[0]
                self.assertEqual(files["integration_terminal_inbound"].name, f"终端入库_{start[:7]}.xlsx")
                self.assertEqual(files["integration_terminal_outbound"].name, f"终端出库_{start[:7]}.xlsx")

    def test_failed_download_does_not_import_stale_file(self):
        with patch.object(materials, "login"), \
                patch.object(materials, "extract_token", return_value="test"), \
                patch.object(materials, "update_cred_json"), \
                patch.object(materials, "export_outbound", return_value=False), \
                patch.object(materials, "finish_exports") as finish:
            with self.assertRaises(RuntimeError):
                materials.collect("2026-08-01", "2026-08-31", output_dir=self.root)
            finish.assert_not_called()

    def test_all_collectors_respect_storage_modes(self):
        for module in (eoms, lines, materials):
            for mode in ("file", "database", "both"):
                with self.subTest(module=module.__name__, mode=mode):
                    folder = self.root / module.__name__ / mode
                    database = folder.parent / f"{mode}.db"
                    downloaded = []
                    def make_file(directory, name, header):
                        directory = Path(directory)
                        directory.mkdir(parents=True, exist_ok=True)
                        path = directory / name
                        wb = Workbook()
                        wb.active.append([header])
                        wb.active.append(["test-1"])
                        wb.save(path)
                        wb.close()
                        downloaded.append(path)
                        return path
                    with ExitStack() as stack:
                        if module is eoms:
                            stack.enter_context(patch.object(eoms, "export_qiwan", side_effect=lambda *a, **kw: make_file(kw["output_dir"], "服务类产品支撑工单_2026-08.xlsx", "工单号")))
                        elif module is lines:
                            stack.enter_context(patch.object(lines, "load_token", return_value=("test", "test")))
                            stack.enter_context(patch.object(lines, "export_order_data", side_effect=lambda *a: make_file(a[3], "一体化专线拆机清单_2026-08.xlsx", "工单号")))
                        else:
                            stack.enter_context(patch.object(materials, "login"))
                            stack.enter_context(patch.object(materials, "extract_token", return_value="test"))
                            stack.enter_context(patch.object(materials, "update_cred_json"))
                            stack.enter_context(patch.object(materials, "export_outbound", side_effect=lambda *a: bool(make_file(a[-1], "终端出库_2026-08.xlsx", "物料名称"))))
                            stack.enter_context(patch.object(materials, "export_inbound", side_effect=lambda *a: bool(make_file(a[-1], "终端入库_2026-08.xlsx", "物料名称"))))
                            stack.enter_context(patch.object(materials, "export_material_baseline", side_effect=lambda *a: bool(make_file(a[-1], "全省物资基准库.xlsx", "物料名称"))))
                        result = module.collect("2026-08-01", "2026-08-31", mode=mode,
                                                output_dir=folder, database=database)
                        if mode != "file":
                            before_count = len(downloaded)
                            cached = module.collect("2026-08-01", "2026-08-31", mode=mode,
                                                    output_dir=folder, database=database)
                            self.assertTrue(cached["skipped"])
                            self.assertEqual(len(downloaded), before_count)
                            if mode == "database":
                                module.collect("2026-08-01", "2026-08-31", mode=mode,
                                               output_dir=folder, database=database, refresh=True)
                                self.assertGreater(len(downloaded), before_count)
                    self.assertEqual(database.exists(), mode != "file")
                    self.assertFalse((database.parent / "raw").exists())
                    if mode == "database":
                        self.assertEqual(result["files"], {})
                        self.assertFalse(folder.exists())
                        self.assertTrue(all(not path.exists() for path in downloaded))
                    else:
                        self.assertTrue(all(path.exists() for path in downloaded))

    def test_reuse_existing_database_mode_preserves_user_files(self):
        path = self.workbook("一体化专线拆机清单_2026-08.xlsx", ["工单号"], [["order-1"]])
        with patch.object(lines, "load_token") as token:
            result = lines.collect("2026-08-01", "2026-08-31", mode="database",
                                   reuse_existing=True, output_dir=self.root, database=self.database)
        token.assert_not_called()
        self.assertTrue(path.exists())
        self.assertEqual(result["files"], {})

    def test_file_mode_does_not_create_database(self):
        path = self.workbook("orders.xlsx", ["工单号"], [["order-1"]])
        for module in (eoms, lines, materials):
            module.finish_exports({"integration_removal_order": path}, "2026-08-01",
                                  "2026-08-31", "file", self.database)
        self.assertFalse(self.database.exists())

    def test_empty_material_export_imports_zero_rows(self):
        code = "integration_material_baseline"
        path = self.workbook("baseline.xlsx", ["物料名称"], [])
        for module in (eoms, lines, materials):
            with self.subTest(module=module.__name__):
                result = module.finish_exports({code: path}, "2026-08-01", "2026-08-31",
                                               "both", self.database)["etl"][code]
                self.assertEqual(result["read"], 0)
                self.assertEqual(result["failed"], 0)

    def test_missing_key_is_reported(self):
        path = self.workbook("orders.xlsx", ["工单号", "名称"], [[None, "missing"]])
        for module in (eoms, lines, materials):
            with self.subTest(module=module.__name__), self.assertRaises(RuntimeError):
                module.finish_exports({"integration_removal_order": path}, "2026-08-01",
                                      "2026-08-31", "both", self.database)

    def test_invalid_dates_rejected_before_download(self):
        for start, end in [("2026-08-32", "2026-08-31"),
                           ("2026-08-02", "2026-08-01"),
                           ("2026-08-01", "2026-09-01"),
                           ("20260801", "2026-08-31")]:
            for module in (eoms, lines, materials):
                with self.subTest(module=module.__name__, start=start, end=end):
                    with patch.object(module, "covered_result") as coverage:
                        with self.assertRaises(ValueError):
                            module.collect(start, end, database=self.database, output_dir=self.root)
                        coverage.assert_not_called()
            with self.assertRaises(ValueError):
                validate_dates(start, end)

    def test_collectors_keep_uniform_cli_defaults(self):
        for module in (eoms, lines, materials):
            with self.subTest(module=module.__name__):
                cli = module.parser("test")
                args = cli.parse_args(["--start-date", "2026-08-01", "--end-date", "2026-08-31"])
                self.assertEqual(args.mode, "both")
                self.assertEqual(args.database, Path("data/quality_assessment.db"))
                self.assertEqual(args.output_dir, Path(r"D:\edge_download"))
                self.assertFalse(args.reuse_existing)
                self.assertFalse(args.refresh)
                args = cli.parse_args(["--start-date", "2026-08-01", "--end-date", "2026-08-31",
                                       "--mode", "database", "--database", str(self.database),
                                       "--output-dir", str(self.root), "--reuse-existing", "--refresh"])
                self.assertEqual(args.mode, "database")
                self.assertEqual(args.database, self.database)
                self.assertEqual(args.output_dir, self.root)
                self.assertTrue(args.reuse_existing)
                self.assertTrue(args.refresh)

    def test_metrics_missing_source_fails_before_output(self):
        with self.assertRaises(FileNotFoundError):
            export("2026-08-01", "2026-08-31", self.root, self.root / "out")
        self.assertFalse((self.root / "out").exists())


if __name__ == "__main__":
    unittest.main()
