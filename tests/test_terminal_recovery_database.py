import gc
import json
import sqlite3
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from storage.terminal_recovery import REPORT_SHEETS, SOURCE_FILES, restore_sources, save_source
from metrics.terminal_recovery import terminal_recovery_export as metrics
from metrics.terminal_recovery.report import read_report
from metrics.installation.common import save_metric


class TerminalDatabaseTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.database = self.root / "test.db"

    def tearDown(self):
        gc.collect()
        self.temp.cleanup()

    def seed(self, start="2026-08-01", end="2026-08-31"):
        wb = Workbook()
        wb.active.append(["identifier", "quantity", "date"])
        wb.active.append(["0001", 2, datetime(2026, 8, 15)])
        wb.active.append(["0001", 2, datetime(2026, 8, 15)])
        wb.active["A1"].font = Font(bold=True)
        source = self.root / "source.xlsx"
        wb.save(source)
        wb.close()
        result = {code: save_source(self.database, code, source, start, end, "test") for code in SOURCE_FILES}
        source.unlink()
        return result

    def test_sources_are_restored_from_database_without_original_files(self):
        expected = self.seed()
        output = self.root / "restored"
        self.assertEqual(restore_sources(self.database, "2026-08-01", "2026-08-31", output), expected)
        for name in SOURCE_FILES.values():
            wb = load_workbook(output / name.format(month="2026-08"))
            self.assertEqual(wb.active.max_row, 3)
            self.assertEqual(wb.active["A2"].value, "0001")
            self.assertEqual(wb.active["B3"].value, 2)
            self.assertEqual(wb.active["C2"].value, datetime(2026, 8, 15))
            self.assertTrue(wb.active["A1"].font.bold)
            wb.close()

    def test_other_period_is_not_silently_used(self):
        self.seed()
        with self.assertRaises(ValueError):
            restore_sources(self.database, "2026-09-01", "2026-09-30", self.root / "missing")
        self.assertFalse((self.root / "missing").exists())

    def test_restoration_uses_sqlite_row_values(self):
        ids = self.seed()
        code = next(iter(ids))
        conn = sqlite3.connect(self.database)
        with conn:
            conn.execute("UPDATE tr_source_row SET cells_json=? WHERE snapshot_id=? AND row_number=2",
                         (json.dumps(["changed-in-db", 9, None]), ids[code]))
        conn.close()
        output = self.root / "restored"
        restore_sources(self.database, "2026-08-01", "2026-08-31", output)
        wb = load_workbook(output / SOURCE_FILES[code].format(month="2026-08"))
        self.assertEqual(wb.active["A2"].value, "changed-in-db")
        self.assertEqual(wb.active["B2"].value, 9)
        wb.close()

    def test_missing_snapshot_rows_are_rejected(self):
        ids = self.seed()
        conn = sqlite3.connect(self.database)
        with conn:
            conn.execute("DELETE FROM tr_source_row WHERE snapshot_id=? AND row_number=2", (next(iter(ids.values())),))
        conn.close()
        with self.assertRaises(ValueError):
            restore_sources(self.database, "2026-08-01", "2026-08-31", self.root / "bad")

    def report(self, bad_formula=False):
        wb = Workbook()
        wb.remove(wb.active)
        for name in REPORT_SHEETS:
            ws = wb.create_sheet(name)
            if name == "按地市回收率汇总":
                ws.append(["地市", "B", "C", "D", "E", "应拆回设备数", "已拆回设备数量", "终端回收率"])
                ws.append(["杭州", 1, 2, 3, 4, 10, 8, "=IFERROR(G2/F2,0)"])
            elif name == "sheet1":
                ws.append(["型号匹配", "业务"])
                ws.append(["ONU", 8])
                ws.append([None, None])
                ws.append(["型号匹配", "地市"])
                ws.append(["ONU", "=1+1" if bad_formula else 8])
            else:
                ws.append(["header"])
                ws.append(["value"])
        path = self.root / "report.xlsx"
        wb.save(path)
        wb.close()
        return path

    def test_two_sheets_use_standard_metric_tables(self):
        report = read_report(self.report(), "2026-08-01", "2026-08-31", {}, [])
        self.assertEqual(set(report["sheets"]), {"sheet1", "按地市回收率汇总"})
        self.assertEqual(report["sheets"]["按地市回收率汇总"]["rows"][0]["终端回收率"], 0.8)
        ids = [save_metric(self.database, report, []) for _ in range(2)]
        self.assertNotEqual(*ids)
        conn = sqlite3.connect(self.database)
        try:
            tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            self.assertNotIn("tr_report", tables)
            self.assertNotIn("tr_report_sheet", tables)
            self.assertEqual(conn.execute("SELECT count(*) FROM ads_metric_result").fetchone()[0], 18)
            self.assertEqual(conn.execute("SELECT numerator,denominator,metric_value FROM ads_metric_result WHERE metric_run_id=? AND metric_code='terminal_recovery_rate'", (ids[0],)).fetchone(), (8, 10, 0.8))
        finally:
            conn.close()

    def test_unexpected_formula_is_rejected(self):
        with self.assertRaises(ValueError):
            read_report(self.report(True), "2026-08-01", "2026-08-31", {}, [])

    def test_metrics_modes_select_sources_and_destinations(self):
        self.seed()
        template = self.report()
        for mode in ("file", "database", "both"):
            with self.subTest(mode=mode):
                output_dir = self.root / mode
                def calculate(start, end, source, target):
                    self.assertTrue((Path(source) / "物料名称表.xlsx").is_file())
                    Path(target).mkdir(parents=True, exist_ok=True)
                    target = Path(target) / "终端回收率汇总表_2026-08.xlsx"
                    shutil.copy2(template, target)
                    return target, {}
                with patch.object(metrics, "export", side_effect=calculate), \
                        patch("storage.terminal_recovery.restore_sources", wraps=restore_sources) as restore:
                    path, stats = metrics.run_export("2026-08-01", "2026-08-31", mode=mode,
                                                    database=self.database, output_dir=output_dir)
                self.assertEqual(restore.call_count, 1)
                self.assertFalse(Path(restore.call_args.args[3]).exists())
                conn = sqlite3.connect(self.database)
                try:
                    if mode == "file":
                        self.assertIsNone(stats["metric_run_id"])
                        self.assertIsNone(conn.execute("SELECT name FROM sqlite_master WHERE name='metric_run'").fetchone())
                    else:
                        self.assertEqual(conn.execute("SELECT count(*) FROM ads_metric_result WHERE metric_run_id=?", (stats["metric_run_id"],)).fetchone()[0], 9)
                finally:
                    conn.close()
                if mode == "database":
                    self.assertIsNone(path)
                    self.assertIsNone(stats["json_file"])
                    self.assertFalse(output_dir.exists())
                else:
                    self.assertTrue(path.is_file())
                    data = json.loads(Path(stats["json_file"]).read_text(encoding="utf-8"))
                    self.assertEqual(set(data["sheets"]), {"sheet1", "按地市回收率汇总"})
                    self.assertEqual(data["metric_run_id"], stats["metric_run_id"])
                    self.assertEqual(data["sheets"]["sheet1"]["business"]["rows"][0]["业务"], 8)


if __name__ == "__main__":
    unittest.main()
