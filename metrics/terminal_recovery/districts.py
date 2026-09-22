"""District cleaning, auditable denominators and recovery summaries."""

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

CITIES = set("杭州 嘉兴 宁波 温州 金华 绍兴 湖州 台州 衢州 丽水 舟山".split())
EXCLUDED = {"集团客户中心", "电子商务中心"}
EXCLUDED_LINE_TYPES = {"行业视频行业版平台基础", "行业视频-小微版"}
INCLUDED_SERVICE_CATEGORIES = {"E企组网", "安全终端部署服务"}
MODELS = ("ONU", "FTTO", "摄像头", "WIFI路由器", "专线卫士")
NON_EQI_VIDEO_COLUMN = "拆机工单数（剔除E企组网和视频监控）"
VIDEO_ORDER_COLUMN = "视频监控拆机工单数"
EQI_DEVICE_COLUMN = "E企组网应拆回设备数"
EXPECTED_COLUMN = "应拆回设备数"
RECOVERED_COLUMN = "已拆回设备数量"
RATE_COLUMN = "终端回收率"


def text(value):
    return "" if value is None else str(value).strip()


def city(value):
    return text(value).removesuffix("市")


def records(path, sheet=None):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [text(v) for v in next(rows)]
        for number, values in enumerate(rows, 2):
            if any(v is not None for v in values):
                yield number, dict(zip(headers, values))
    finally:
        wb.close()


def load_mapping(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        headers = [text(v) for v in next(wb.active.iter_rows(values_only=True))]
        if not {"地市", "区县", "标准区县"}.issubset(headers):
            raise ValueError("区县映射表必须包含地市、区县、标准区县三列")
    finally:
        wb.close()
    return DistrictMap(row for _, row in records(path))


class DistrictMap:
    def __init__(self, rows):
        self.entries = set()
        for row in rows:
            municipality = city(row.get("地市"))
            standard = text(row.get("标准区县"))
            if municipality not in CITIES or not standard:
                raise ValueError(f"区县映射缺少有效地市或标准区县：{row}")
            for alias in (text(row.get("区县")), standard):
                if alias and city(alias) not in CITIES:
                    self.entries.add((municipality, alias, standard))
        if not self.entries:
            raise ValueError("区县映射表为空")

    def match(self, value, municipality, keyword=False):
        value = text(value)
        if not value or city(value) in CITIES:
            return ""
        candidates = [(c, a, s) for c, a, s in self.entries
                      if (a in value if keyword else a == value)]
        local = [item for item in candidates if item[0] == city(municipality)]
        if local:
            candidates = local
        # A known city must not acquire a district belonging to another city.
        elif city(municipality) in CITIES:
            return ""
        if not candidates:
            return ""
        longest = max(len(a) for _, a, _ in candidates)
        matches = {(c, s) for c, a, s in candidates if len(a) == longest}
        return next(iter(matches))[1] if len(matches) == 1 else ""


def clean_row(row, mapping, service=False):
    original = text(row.get("安装区县" if service else "区县"))
    municipality = city(row.get("安装地市" if service else "地市"))
    if service and not municipality:
        municipality = city(row.get("所属地市"))
    if original in EXCLUDED:
        return municipality, "", "剔除中心", False
    if not service and not original and row.get("业务类型") in EXCLUDED_LINE_TYPES:
        return municipality, "", "空区县业务剔除", False
    district = mapping.match(original, municipality)
    if district:
        return municipality, district, "区县直接映射", True
    if service and municipality and municipality == city(row.get("所属地市")):
        district = mapping.match(row.get("所属区县"), municipality)
        if district:
            return municipality, district, "所属区县映射", True
    for field in (("主题", "客户机房地址") if service else ("工单主题",)):
        district = mapping.match(row.get(field), municipality, keyword=True)
        if district:
            return municipality, district, field + "关键词", True
    return municipality, "", "未匹配或歧义", True


def resolve_mapping(output_dir, explicit=None):
    if explicit:
        path = Path(explicit)
        if not path.is_file():
            raise FileNotFoundError(path)
        return path
    default = Path(__file__).resolve().parent / "区县信息汇总.xlsx"
    if not default.is_file():
        raise FileNotFoundError(f"默认区县映射表不存在：{default}")
    return default


def build_district_report(line_source, service_source, recovery_file, mapping_file, output_file):
    mapping = load_mapping(mapping_file)
    non_eqi_video, video_orders, eqi_devices, recovered = Counter(), Counter(), Counter(), Counter()
    models = defaultdict(Counter)
    wb = Workbook(write_only=True)
    summary = wb.create_sheet("区县汇总")
    bottom = wb.create_sheet("回收率后十位")
    details = wb.create_sheet("拆机区县清洗明细")
    details.append(["来源", "源行号", "工单号", "地市", "原区县", "标准区县", "处理依据", "是否保留", "是否计入区县分母", "主题", "客户机房地址"])
    counts = Counter()
    for source, service in ((line_source, False), (service_source, True)):
        for number, row in records(source):
            municipality, district, reason, keep = clean_row(row, mapping, service)
            district_valid = keep and municipality in CITIES and bool(district)
            business_type = text(row.get("服务类别" if service else "业务类型"))
            in_calculation_scope = (
                business_type in INCLUDED_SERVICE_CATEGORIES if service
                else business_type not in EXCLUDED_LINE_TYPES
            )
            included = district_valid and in_calculation_scope
            if included:
                key = municipality, district
                if business_type == "视频监控":
                    video_orders[key] += 1
                elif business_type != "E企组网":
                    non_eqi_video[key] += 1
                if service:
                    eqi_devices[key] += sum(float(row.get(column) or 0)
                                            for column in ("路由器数量", "FTTR数量", "光AP数量"))
            counts[("服务类" if service else "一体化") + ":" + reason] += 1
            if district_valid and not in_calculation_scope:
                counts[("服务类" if service else "一体化") + ":统计口径剔除"] += 1
            details.append([Path(source).name, number, row.get("工单号"), municipality,
                            row.get("安装区县" if service else "区县"), district, reason,
                            int(keep), int(included), row.get("主题" if service else "工单主题"), row.get("客户机房地址")])
    returns = wb.create_sheet("回收区县明细")
    returns.append(["源行号", "关联单据号", "地市", "原区县", "统计区县", "型号匹配", "总数量", "计入区县分子"])
    for number, row in records(recovery_file, "匹配后拆回设备清单（三类标签已去重）"):
        municipality, original = city(row.get("地市")), text(row.get("区县"))
        district = "开发区" if original == "南城区" else original
        district = mapping.match(district, municipality)
        quantity = float(row.get("总数量") or 0)
        included = municipality in CITIES and bool(district) and district not in EXCLUDED and city(district) not in CITIES
        if included:
            recovered[municipality, district] += quantity
            models[municipality, district][text(row.get("型号匹配"))] += quantity
        returns.append([number, row.get("关联单据号"), municipality, original, district, row.get("型号匹配"), quantity, int(included)])
    rows = []
    all_keys = non_eqi_video.keys() | video_orders.keys() | eqi_devices.keys() | recovered.keys()
    for municipality, district in sorted(all_keys):
        key = municipality, district
        expected = non_eqi_video[key] + video_orders[key] * 2 + eqi_devices[key]
        numerator = recovered[key]
        rows.append({"地市": municipality, "区县": district,
                     NON_EQI_VIDEO_COLUMN: non_eqi_video[key],
                     VIDEO_ORDER_COLUMN: video_orders[key], EQI_DEVICE_COLUMN: eqi_devices[key],
                     EXPECTED_COLUMN: expected, RECOVERED_COLUMN: numerator,
                     RATE_COLUMN: numerator / expected if expected else None,
                     **{model: models[key][model] for model in MODELS}})
    ranked = sorted((r for r in rows if r[EXPECTED_COLUMN] > 0),
                    key=lambda r: (r[RATE_COLUMN], r["地市"], r["区县"]))[:10]
    headers = ["地市", "区县", NON_EQI_VIDEO_COLUMN, VIDEO_ORDER_COLUMN,
               EQI_DEVICE_COLUMN, EXPECTED_COLUMN, RECOVERED_COLUMN, RATE_COLUMN, *MODELS]
    for ws, data in ((summary, rows), (bottom, ranked)):
        ws.freeze_panes = "C2"
        for column in ("A", "B"):
            ws.column_dimensions[column].width = 16
        for column in ("C", "D", "E", "F", "G", "H"):
            ws.column_dimensions[column].width = 25
        ws.append(headers)
        for row in data:
            values = [row[h] for h in headers]
            cell = WriteOnlyCell(ws, value=row[RATE_COLUMN])
            cell.number_format = "0.00%"
            values[7] = cell
            ws.append(values)
    ws = wb.create_sheet("清洗统计")
    ws.append(["处理依据", "记录数"])
    for key, count in sorted(counts.items()):
        ws.append([key, count])
    wb.save(output_file)
    return {"mapping_file": str(Path(mapping_file).resolve()), "audit_file": str(Path(output_file).resolve()),
            "rows": rows, "bottom10": ranked, "cleaning_counts": dict(counts)}


def metric_rows(report):
    result = []
    for row in report["rows"]:
        for column in (NON_EQI_VIDEO_COLUMN, VIDEO_ORDER_COLUMN, EQI_DEVICE_COLUMN,
                       EXPECTED_COLUMN, RECOVERED_COLUMN, RATE_COLUMN, *MODELS):
            if row[column] is None:
                continue
            rate = column == RATE_COLUMN
            result.append({"metric_code": "terminal_recovery_rate" if rate else "terminal_recovery_count",
                           "dimension_type": "district_summary",
                           "dimension": {"sheet": "区县汇总", "table": "district_summary",
                                         "label_column": "地市·标准区县",
                                         "label": row["地市"] + "·" + row["区县"], "column": column},
                           "numerator": row[RECOVERED_COLUMN] if rate else row[column],
                           "denominator": row[EXPECTED_COLUMN] if rate else 1,
                           "metric_value": row[column]})
    return result
