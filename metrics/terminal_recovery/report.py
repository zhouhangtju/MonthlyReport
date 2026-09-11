"""Convert the two public recovery sheets to JSON and standard metric rows."""

import math
from openpyxl import load_workbook

METRIC_CODE = "terminal_recovery"
METRIC_VERSION = "2.0"
CITY_SHEET = "按地市回收率汇总"
DETAIL_SHEET = "sheet1"


def number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Expected a finite numeric result, got {value!r}")
    return value


def table(rows, header_index, end_index):
    header = list(rows[header_index])
    while header and header[-1] is None:
        header.pop()
    if not header or any(not isinstance(c, str) or not c for c in header) or len(set(header)) != len(header):
        raise ValueError("Invalid summary headers")
    records = []
    for row in rows[header_index + 1:end_index]:
        if all(v is None for v in row):
            continue
        if not row[0]:
            raise ValueError("Summary row missing label")
        records.append(dict(zip(header, row[:len(header)])))
    return {"columns": header, "rows": records}


def read_report(path, start, end, snapshots, source_runs):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        city_rows = list(wb[CITY_SHEET].iter_rows(values_only=True))
        city = table(city_rows, 0, len(city_rows))
        required = {"地市", "应拆回设备数", "已拆回设备数量", "终端回收率"}
        if not required.issubset(city["columns"]):
            raise ValueError("City summary missing required columns")
        for row_number, row in enumerate(city["rows"], 2):
            formula = row["终端回收率"]
            if formula != f"=IFERROR(G{row_number}/F{row_number},0)":
                raise ValueError(f"Unexpected recovery formula: {formula}")
            numerator, denominator = number(row["已拆回设备数量"]), number(row["应拆回设备数"])
            row["终端回收率"] = numerator / denominator if denominator else 0
        detail_rows = list(wb[DETAIL_SHEET].iter_rows(values_only=True))
        headers = [i for i, row in enumerate(detail_rows) if row[0] == "型号匹配"]
        if len(headers) != 2:
            raise ValueError("sheet1 must contain exactly two cross tables")
        business = table(detail_rows, headers[0], headers[1])
        cities = table(detail_rows, headers[1], len(detail_rows))
    finally:
        wb.close()
    sheets = {CITY_SHEET: city, DETAIL_SHEET: {"business": business, "city": cities}}
    results = []
    for sheet, name, data in [(CITY_SHEET, "city_summary", city),
                              (DETAIL_SHEET, "business", business), (DETAIL_SHEET, "city", cities)]:
        for row_index, row in enumerate(data["rows"], 1):
            label_column = data["columns"][0]
            for column_index, column in enumerate(data["columns"][1:], 2):
                value = number(row[column])
                is_rate = sheet == CITY_SHEET and column == "终端回收率"
                results.append({
                    "metric_code": "terminal_recovery_rate" if is_rate else "terminal_recovery_count",
                    "dimension_type": name,
                    "dimension": {"sheet": sheet, "table": name, "row_index": row_index,
                                  "column_index": column_index, "label_column": label_column,
                                  "label": row[label_column], "column": column},
                    "numerator": row["已拆回设备数量"] if is_rate else value,
                    "denominator": row["应拆回设备数"] if is_rate else 1,
                    "metric_value": value,
                })
    return {"metric_code": METRIC_CODE, "metric_version": METRIC_VERSION,
            "period_start": start, "period_end": end,
            "source_runs": source_runs, "source_snapshots": snapshots,
            "sheets": sheets, "results": results}
