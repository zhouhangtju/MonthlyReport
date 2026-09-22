#!/usr/bin/env python3
"""线上生成终端回收率汇总表，重复关联单据号保留首次记录。"""

import argparse
import re
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

# ============================================================
# 配置：第一步——生成基础拆机清单及按地市回收率汇总
# ============================================================

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = Path(r"D:\edge_download")
DEFAULT_OUTPUT_DIR = SCRIPT_DIR.parents[1] / "outputs"
MATERIAL_NAME_SOURCE_FILE = SCRIPT_DIR / "物料名称表.xlsx"

PROVINCE_MATERIAL_SOURCE_FILE = Path("全省物资基准库.xlsx")

LINE_SHEET_NAME = "一体化专线拆机清单（不含小微版）"

BUSINESS_TYPE_COLUMN = "业务类型"

EXCLUDED_TYPES = {"行业视频行业版平台基础", "行业视频-小微版"}

SERVICE_SHEET_NAME = "EOMS服务类拆机清单"

SERVICE_CATEGORY_COLUMN = "服务类别"

INCLUDED_SERVICE_CATEGORIES = {"E企组网", "安全终端部署服务"}

INSTALLATION_CITY_COLUMN = "安装地市"

AFFILIATED_CITY_COLUMN = "所属地市"

AFFILIATED_COUNTY_COLUMN = "所属区县"

RETURN_DEVICE_SHEET_NAME = "一体化拆回设备清单"

OPERATION_TYPE_COLUMN = "操作类型"

REMOVAL_SUMMARY_SHEET_NAME = "拆除清单汇总"

REMOVAL_SUMMARY_COLUMNS = ("工单号", "地市", "区县")

REMOVAL_BUSINESS_TYPE_COLUMN = "业务类型1"

REMOVAL_BUSINESS_TYPE_2_COLUMN = "业务类型2"

VIDEO_BUSINESS_TYPES = {"视频监控", "行业视频MV专线", "行业视频-行业版"}

PASSTHROUGH_BUSINESS_TYPES = {
    "互联网专线",
    "MPLS VPN专线",
    "电路出租",
    "固话",
}

CITY_SUMMARY_SHEET_NAME = "按地市回收率汇总"

CITIES = (
    "金华",
    "温州",
    "湖州",
    "杭州",
    "宁波",
    "丽水",
    "舟山",
    "台州",
    "衢州",
    "嘉兴",
    "绍兴",
)

CITY_SUMMARY_HEADERS = (
    "地市",
    "拆机工单数（不含小微版和行业版平台基础产品）",
    "拆机工单数（剔除E企组网和视频监控）",
    "视频监控拆机工单数",
    "E企组网应拆回设备数",
    "应拆回设备数",
    "已拆回设备数量",
    "终端回收率",
)

ROUTER_QUANTITY_COLUMN = "路由器数量"

FTTR_QUANTITY_COLUMN = "FTTR数量"

LIGHT_AP_QUANTITY_COLUMN = "光AP数量"


# ============================================================
# 配置：第二步——匹配拆回设备、生成去重结果及明细汇总
# ============================================================

RETURN_DEVICE_SHEET = "一体化拆回设备清单"

REMOVAL_SUMMARY_SHEET = "拆除清单汇总"

MATERIAL_NAME_SHEET = "物料名称表"

PROVINCE_MATERIAL_SHEET = "全省物资基准库"

MATCHED_RETURN_SHEET = "匹配后拆回设备清单（三类标签已去重）"

SHEET2_NAME = "sheet2"

SHEET3_NAME = "sheet3"

CITY_SUMMARY_SHEET = "按地市回收率汇总"

DETAIL_SUMMARY_SHEET = "sheet1"

RELATED_ORDER_COLUMN = "关联单据号"

ORDER_NUMBER_COLUMN = "工单号"

BUSINESS_TYPE_1_COLUMN = "业务类型1"

BUSINESS_TYPE_2_COLUMN = "业务类型2"

MATERIAL_NAME_COLUMN = "物料名称"

MATERIAL_ID_COLUMN = "物料ID"

DIVISION_TYPE_COLUMN = "划分类型"

CATEGORY_COLUMN = "分类"

CITY_COLUMN = "地市"

TOTAL_QUANTITY_COLUMN = "总数量"

RECOVERED_DEVICE_COUNT_COLUMN = "已拆回设备数量"

EXPECTED_DEVICE_COUNT_COLUMN = "应拆回设备数"

RECOVERY_RATE_COLUMN = "终端回收率"

STANDARD_CITIES = (
    "金华",
    "温州",
    "湖州",
    "杭州",
    "宁波",
    "丽水",
    "舟山",
    "台州",
    "衢州",
    "嘉兴",
    "绍兴",
)

ORDER_MATCH_COLUMN = "关联单据号匹配"

BUSINESS_CATEGORY_1_COLUMN = "业务类别1"

BUSINESS_CATEGORY_2_COLUMN = "业务类别2"

MATERIAL_NAME_MATCH_COLUMN = "物料名称匹配"

MODEL_MATCH_COLUMN = "型号匹配"

SHEET1_MATCH_COLUMN = "sheet1关联单据号匹配"

NEW_COLUMNS = (
    ORDER_MATCH_COLUMN,
    BUSINESS_CATEGORY_1_COLUMN,
    BUSINESS_CATEGORY_2_COLUMN,
    MATERIAL_NAME_MATCH_COLUMN,
    MODEL_MATCH_COLUMN,
    CATEGORY_COLUMN,
)

ZERO_DIVISION_TYPE_MATERIALS = {
    "共用设备拆回标签",
    "在网业务标签",
    "无设备标签",
}

BASE_EXCLUDED_BUSINESS_TYPES = {
    "E企组网",
    "视频监控",
    "安全终端部署服务",
}

SERVICE_BUSINESS_TYPES = {"E企组网", "安全终端部署服务"}

SERVICE_MODEL_TYPES = {"WIFI路由器", "FTTO", "专线卫士"}

SUMMARY_MODEL_TYPES = ("ONU", "FTTO", "摄像头", "WIFI路由器", "专线卫士")

SUMMARY_BUSINESS_TYPES = (
    "MPLS VPN专线",
    "电路出租",
    "服务类",
    "固话",
    "互联网专线",
    "其他专线",
    "视频监控",
)

SUMMARY_CITIES = (
    ("杭州市", "杭州"),
    ("湖州", "湖州"),
    ("嘉兴", "嘉兴"),
    ("金华", "金华"),
    ("丽水", "丽水"),
    ("宁波", "宁波"),
    ("衢州", "衢州"),
    ("绍兴", "绍兴"),
    ("台州", "台州"),
    ("温州", "温州"),
    ("舟山", "舟山"),
)

PROGRESS_INTERVAL = 5000




# ============================================================
# 公共工具函数
# ============================================================

def is_blank(value):
    """判断单元格是否为空。"""
    return value is None or (isinstance(value, str) and value.strip() == "")

def normalize_header(value):
    """清理表头前后空格。"""
    return value.strip() if isinstance(value, str) else value

def get_column_indexes(ws, column_names):
    """根据第一行表头获取列号。"""
    indexes = {
        normalize_header(cell.value): cell.column
        for cell in ws[1]
        if not is_blank(cell.value)
    }
    missing = [name for name in column_names if name not in indexes]
    if missing:
        raise ValueError(
            f"工作表 {ws.title} 中未找到列：{', '.join(missing)}"
        )
    return {name: indexes[name] for name in column_names}

def normalize_city(value):
    """将地市统一为11个标准名称，温岭归入台州。"""
    city = normalize_text(value)
    if "温岭" in city:
        return "台州"
    for standard_city in STANDARD_CITIES:
        if standard_city in city:
            return standard_city
    return city

def to_number(value):
    """将总数量转换为数值，空值或非数字按0处理。"""
    if is_blank(value):
        return 0
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).strip().replace(",", ""))
    except ValueError:
        return 0

def copy_style(source_cell, target_cell):
    """跨工作簿安全复制单元格样式，不直接复用内部样式编号。"""
    if not source_cell.has_style:
        return

    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format

def copy_cell(source_cell, target_cell):
    """复制单元格的值、样式、超链接和批注。"""
    target_cell.value = source_cell.value
    copy_style(source_cell, target_cell)
    if source_cell.hyperlink:
        target_cell._hyperlink = copy(source_cell.hyperlink)
    if source_cell.comment:
        target_cell.comment = copy(source_cell.comment)


def copy_sheet_dimensions(source_ws, target_ws, copy_rows=True):
    """安全复制列宽和行高，不复制跨工作簿无效的样式编号。"""
    for key, source_dimension in source_ws.column_dimensions.items():
        target_dimension = target_ws.column_dimensions[key]
        for attribute in (
            "width",
            "min",
            "max",
            "bestFit",
            "hidden",
            "outlineLevel",
            "collapsed",
        ):
            setattr(
                target_dimension,
                attribute,
                copy(getattr(source_dimension, attribute)),
            )

    if not copy_rows:
        return

    for row_index, source_dimension in source_ws.row_dimensions.items():
        target_dimension = target_ws.row_dimensions[row_index]
        for attribute in (
            "height",
            "hidden",
            "outlineLevel",
            "collapsed",
            "thickTop",
            "thickBot",
        ):
            setattr(
                target_dimension,
                attribute,
                copy(getattr(source_dimension, attribute)),
            )


def validate_month(month):
    """校验月份参数是否为YYYY-MM格式。"""
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise ValueError(
            f"月份参数格式错误：{month}，正确格式示例为 2026-08"
        )


def get_output_file(month):
    """根据月份参数生成汇总表文件名。"""
    validate_month(month)
    return Path(f"终端回收率汇总表_{month}.xlsx")


def get_business_source_files(month):
    """根据月份参数生成三张业务源表的文件名。"""
    validate_month(month)
    return (
        Path(f"一体化专线拆机清单_{month}.xlsx"),
        Path(f"服务类产品支撑工单_{month}.xlsx"),
        Path(f"终端入库_{month}.xlsx"),
    )


def copy_source_sheet(target_wb, source_file, sheet_name, sheet_index):
    """将源文件活动工作表完整复制到目标工作簿。"""
    if not source_file.exists():
        raise FileNotFoundError(f"当前目录下未找到文件：{source_file}")

    source_wb = load_workbook(source_file, data_only=False)
    try:
        source_ws = source_wb.active
        if sheet_name in target_wb.sheetnames:
            target_wb.remove(target_wb[sheet_name])
        target_ws = target_wb.create_sheet(sheet_name, sheet_index)

        copy_sheet_dimensions(source_ws, target_ws)

        for source_row in source_ws.iter_rows():
            for source_cell in source_row:
                copy_cell(
                    source_cell,
                    target_ws.cell(source_cell.row, source_cell.column),
                )

        for merged_range in source_ws.merged_cells.ranges:
            target_ws.merge_cells(str(merged_range))

        target_ws.freeze_panes = source_ws.freeze_panes
        target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
        target_ws.sheet_format = copy(source_ws.sheet_format)
        if source_ws.auto_filter.ref:
            target_ws.auto_filter.ref = source_ws.auto_filter.ref
    finally:
        source_wb.close()


def create_summary_workbook(
    material_name_source_file=MATERIAL_NAME_SOURCE_FILE,
    province_material_source_file=PROVINCE_MATERIAL_SOURCE_FILE,
):
    """新建汇总工作簿，并写入物料名称表和全省物资基准库。"""
    summary_wb = Workbook()
    summary_wb.remove(summary_wb.active)
    copy_source_sheet(summary_wb, material_name_source_file, MATERIAL_NAME_SHEET, 0)
    copy_source_sheet(
        summary_wb,
        province_material_source_file,
        PROVINCE_MATERIAL_SHEET,
        1,
    )
    return summary_wb


# ============================================================
# 第一步：生成基础工作表
# ============================================================

def get_header_indexes(ws):
    """返回清理过前后空格的表头与列号映射。"""
    return {
        normalize_header(cell.value): cell.column
        for cell in ws[1]
        if not is_blank(cell.value)
    }

def write_filtered_sheet(
    summary_wb,
    source_file,
    sheet_name,
    sheet_index,
    filter_column,
    include_values=None,
    exclude_values=None,
    remove_blank=False,
    fill_blank_from=None,
    city_columns=(),
):
    """筛选源表数据，并复制到汇总表的指定位置。"""
    source_wb = load_workbook(source_file, data_only=False)
    source_ws = source_wb.active

    header_indexes = get_header_indexes(source_ws)
    if filter_column not in header_indexes:
        source_wb.close()
        raise ValueError(f"{source_file} 中未找到列：{filter_column}")
    filter_col = header_indexes[filter_column]

    fill_blank_from = fill_blank_from or {}
    fill_columns = set(fill_blank_from) | set(fill_blank_from.values())
    missing_fill_columns = [
        name for name in fill_columns if name not in header_indexes
    ]
    if missing_fill_columns:
        source_wb.close()
        raise ValueError(
            f"{source_file} 中未找到列：{', '.join(missing_fill_columns)}"
        )
    fill_column_indexes = {
        name: header_indexes[name] for name in fill_columns
    }

    missing_city_columns = [
        name for name in city_columns if name not in header_indexes
    ]
    if missing_city_columns:
        source_wb.close()
        raise ValueError(
            f"{source_file} 中未找到列：{', '.join(missing_city_columns)}"
        )
    city_column_indexes = {
        name: header_indexes[name] for name in city_columns
    }

    if sheet_name in summary_wb.sheetnames:
        summary_wb.remove(summary_wb[sheet_name])
    target_ws = summary_wb.create_sheet(sheet_name, sheet_index)

    # 复制源表的列宽和工作表设置。
    copy_sheet_dimensions(source_ws, target_ws, copy_rows=False)
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines

    kept_count = 0
    excluded_count = 0
    target_row = 1

    for source_row in source_ws.iter_rows():
        if source_row[0].row > 1:
            filter_value = source_row[filter_col - 1].value
            should_skip = (
                (include_values is not None and filter_value not in include_values)
                or (exclude_values is not None and filter_value in exclude_values)
                or (remove_blank and is_blank(filter_value))
            )
            if should_skip:
                excluded_count += 1
                continue

        for source_cell in source_row:
            copy_cell(source_cell, target_ws.cell(target_row, source_cell.column))

        # 目标字段为空时，用同一行指定的备用字段补充。
        for target_column, fallback_column in fill_blank_from.items():
            target_cell = target_ws.cell(
                target_row, fill_column_indexes[target_column]
            )
            if is_blank(target_cell.value):
                target_cell.value = source_row[
                    fill_column_indexes[fallback_column] - 1
                ].value

        # 数据写入目标表时立即清洗地市，确保保存后的sheet只保留标准地市名。
        if source_row[0].row > 1:
            for city_column_name, city_column_index in city_column_indexes.items():
                city_cell = target_ws.cell(target_row, city_column_index)
                if is_blank(city_cell.value):
                    continue

                original_city = str(city_cell.value).strip()
                normalized_city = normalize_city(original_city)
                if normalized_city not in CITIES:
                    source_wb.close()
                    raise ValueError(
                        f"{source_file} 第 {source_row[0].row} 行的"
                        f"{city_column_name}无法识别：{original_city}"
                    )
                city_cell.value = normalized_city

        source_height = source_ws.row_dimensions[source_row[0].row].height
        if source_height is not None:
            target_ws.row_dimensions[target_row].height = source_height

        if source_row[0].row > 1:
            kept_count += 1
        target_row += 1

    if source_ws.auto_filter.ref:
        target_ws.auto_filter.ref = f"A1:{target_ws.cell(target_row - 1, source_ws.max_column).coordinate}"

    source_wb.close()
    return kept_count, excluded_count

def map_business_type_2(business_type_1):
    """根据业务类型1生成业务类型2。"""
    if business_type_1 in INCLUDED_SERVICE_CATEGORIES:
        return "服务类"
    if business_type_1 in VIDEO_BUSINESS_TYPES:
        return "视频监控"
    if business_type_1 in PASSTHROUGH_BUSINESS_TYPES:
        return business_type_1
    return "其他专线"

def clean_sheet_cities(ws, city_column_name):
    """将指定地市列统一为11个标准值。"""
    city_column = get_column_indexes(ws, (city_column_name,))[city_column_name]
    unknown_cities = set()
    cleaned_count = 0

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row, city_column)
        if is_blank(cell.value):
            continue

        normalized = normalize_city(cell.value)
        if normalized not in CITIES:
            unknown_cities.add(str(cell.value).strip())
            continue

        cell.value = normalized
        cleaned_count += 1

    if unknown_cities:
        values = "、".join(sorted(unknown_cities))
        raise ValueError(
            f"工作表 {ws.title} 的{city_column_name}列中存在无法识别的地市：{values}"
        )

    return cleaned_count

def build_removal_summary_sheet(summary_wb):
    """合并专线和EOMS拆机清单，生成统一的拆除清单汇总。"""
    source_configs = (
        (
            LINE_SHEET_NAME,
            BUSINESS_TYPE_COLUMN,
            {"工单号": "工单号", "地市": "地市", "区县": "区县"},
        ),
        (
            SERVICE_SHEET_NAME,
            SERVICE_CATEGORY_COLUMN,
            {
                "工单号": "工单号",
                "地市": INSTALLATION_CITY_COLUMN,
                "区县": AFFILIATED_COUNTY_COLUMN,
            },
        ),
    )

    if REMOVAL_SUMMARY_SHEET_NAME in summary_wb.sheetnames:
        summary_wb.remove(summary_wb[REMOVAL_SUMMARY_SHEET_NAME])
    target_ws = summary_wb.create_sheet(REMOVAL_SUMMARY_SHEET_NAME, 5)

    # 使用一体化专线拆机清单的表头样式。
    line_ws = summary_wb[LINE_SHEET_NAME]
    line_columns = get_column_indexes(
        line_ws, (*REMOVAL_SUMMARY_COLUMNS, BUSINESS_TYPE_COLUMN)
    )
    for target_col, column_name in enumerate(REMOVAL_SUMMARY_COLUMNS, start=1):
        copy_cell(line_ws.cell(1, line_columns[column_name]), target_ws.cell(1, target_col))
    copy_cell(
        line_ws.cell(1, line_columns[BUSINESS_TYPE_COLUMN]),
        target_ws.cell(1, 4),
    )
    target_ws.cell(1, 4).value = REMOVAL_BUSINESS_TYPE_COLUMN
    copy_cell(
        line_ws.cell(1, line_columns[BUSINESS_TYPE_COLUMN]),
        target_ws.cell(1, 5),
    )
    target_ws.cell(1, 5).value = REMOVAL_BUSINESS_TYPE_2_COLUMN

    target_row = 2
    source_counts = {}

    for source_sheet_name, type_column, source_column_map in source_configs:
        source_ws = summary_wb[source_sheet_name]
        columns = get_column_indexes(
            source_ws, (*source_column_map.values(), type_column)
        )
        source_count = 0

        for source_row in range(2, source_ws.max_row + 1):
            for target_col, column_name in enumerate(REMOVAL_SUMMARY_COLUMNS, start=1):
                copy_cell(
                    source_ws.cell(
                        source_row, columns[source_column_map[column_name]]
                    ),
                    target_ws.cell(target_row, target_col),
                )

            # 汇总表中的地市统一为11个标准值。
            city_cell = target_ws.cell(target_row, 2)
            if not is_blank(city_cell.value):
                normalized_city = normalize_city(city_cell.value)
                if normalized_city not in CITIES:
                    raise ValueError(
                        f"工作表 {source_sheet_name} 中存在无法识别的地市："
                        f"{city_cell.value}"
                    )
                city_cell.value = normalized_city

            source_type_cell = source_ws.cell(source_row, columns[type_column])
            copy_cell(source_type_cell, target_ws.cell(target_row, 4))
            copy_cell(source_type_cell, target_ws.cell(target_row, 5))
            target_ws.cell(target_row, 5).value = map_business_type_2(
                source_type_cell.value
            )
            target_row += 1
            source_count += 1

        source_counts[source_sheet_name] = source_count

    target_ws.freeze_panes = line_ws.freeze_panes
    target_ws.sheet_view.showGridLines = line_ws.sheet_view.showGridLines
    if target_row > 2:
        target_ws.auto_filter.ref = f"A1:E{target_row - 1}"

    return source_counts

def build_city_recovery_summary_sheet(summary_wb):
    """按固定地市统计拆机工单数及应拆回设备数。"""
    removal_ws = summary_wb[REMOVAL_SUMMARY_SHEET_NAME]
    removal_columns = get_column_indexes(
        removal_ws, ("地市", REMOVAL_BUSINESS_TYPE_COLUMN)
    )

    total_orders = Counter()
    orders_without_eqi_and_video = Counter()
    video_orders = Counter()

    for row in range(2, removal_ws.max_row + 1):
        city = normalize_city(removal_ws.cell(row, removal_columns["地市"]).value)
        if city not in CITIES:
            continue

        business_type = removal_ws.cell(
            row, removal_columns[REMOVAL_BUSINESS_TYPE_COLUMN]
        ).value
        total_orders[city] += 1
        if business_type not in {"E企组网", "视频监控"}:
            orders_without_eqi_and_video[city] += 1
        if business_type == "视频监控":
            video_orders[city] += 1

    service_ws = summary_wb[SERVICE_SHEET_NAME]
    service_columns = get_column_indexes(
        service_ws,
        (
            INSTALLATION_CITY_COLUMN,
            ROUTER_QUANTITY_COLUMN,
            FTTR_QUANTITY_COLUMN,
            LIGHT_AP_QUANTITY_COLUMN,
        ),
    )
    eqi_device_counts = Counter()

    for row in range(2, service_ws.max_row + 1):
        city = normalize_city(
            service_ws.cell(row, service_columns[INSTALLATION_CITY_COLUMN]).value
        )
        if city not in CITIES:
            continue

        eqi_device_counts[city] += sum(
            to_number(service_ws.cell(row, service_columns[column]).value)
            for column in (
                ROUTER_QUANTITY_COLUMN,
                FTTR_QUANTITY_COLUMN,
                LIGHT_AP_QUANTITY_COLUMN,
            )
        )

    if CITY_SUMMARY_SHEET_NAME in summary_wb.sheetnames:
        summary_wb.remove(summary_wb[CITY_SUMMARY_SHEET_NAME])
    target_ws = summary_wb.create_sheet(CITY_SUMMARY_SHEET_NAME, 6)

    for column, header in enumerate(CITY_SUMMARY_HEADERS, start=1):
        target_ws.cell(1, column).value = header

    for row, city in enumerate(CITIES, start=2):
        eqi_count = eqi_device_counts[city]
        if isinstance(eqi_count, float) and eqi_count.is_integer():
            eqi_count = int(eqi_count)

        target_ws.cell(row, 1).value = city
        target_ws.cell(row, 2).value = total_orders[city]
        target_ws.cell(row, 3).value = orders_without_eqi_and_video[city]
        target_ws.cell(row, 4).value = video_orders[city]
        target_ws.cell(row, 5).value = eqi_count
        target_ws.cell(row, 6).value = (
            orders_without_eqi_and_video[city]
            + video_orders[city] * 2
            + eqi_count
        )
        # “已拆回设备数量”和“终端回收率”暂时留空。
        target_ws.cell(row, 7).value = None
        target_ws.cell(row, 8).value = None

    # 按参考图片设置简洁的表格样式。
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    for row in target_ws.iter_rows(min_row=1, max_row=len(CITIES) + 1, min_col=1, max_col=8):
        for cell in row:
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = thin_border

    target_ws.row_dimensions[1].height = 58
    for row in range(2, len(CITIES) + 2):
        target_ws.row_dimensions[row].height = 22
    target_ws.column_dimensions["A"].width = 12
    for column in ("B", "C", "D", "E", "F", "G", "H"):
        target_ws.column_dimensions[column].width = 24
    target_ws.freeze_panes = "A2"
    target_ws.sheet_view.showGridLines = False

    return len(CITIES)

def build_workbook(
    month,
    line_source_file=None,
    service_source_file=None,
    return_device_source_file=None,
    material_name_source_file=MATERIAL_NAME_SOURCE_FILE,
    province_material_source_file=PROVINCE_MATERIAL_SOURCE_FILE,
    output_file=None,
):
    """新建汇总表，依次生成第一至第七个工作表。"""
    default_source_files = get_business_source_files(month)
    line_source_file = line_source_file or default_source_files[0]
    service_source_file = service_source_file or default_source_files[1]
    return_device_source_file = return_device_source_file or default_source_files[2]
    if output_file is None:
        output_file = get_output_file(month)

    summary_wb = create_summary_workbook(
        material_name_source_file=material_name_source_file,
        province_material_source_file=province_material_source_file,
    )

    line_counts = write_filtered_sheet(
        summary_wb=summary_wb,
        source_file=line_source_file,
        sheet_name=LINE_SHEET_NAME,
        sheet_index=2,
        filter_column=BUSINESS_TYPE_COLUMN,
        exclude_values=EXCLUDED_TYPES,
        city_columns=("地市",),
    )
    service_counts = write_filtered_sheet(
        summary_wb=summary_wb,
        source_file=service_source_file,
        sheet_name=SERVICE_SHEET_NAME,
        sheet_index=3,
        filter_column=SERVICE_CATEGORY_COLUMN,
        include_values=INCLUDED_SERVICE_CATEGORIES,
        fill_blank_from={INSTALLATION_CITY_COLUMN: AFFILIATED_CITY_COLUMN},
        city_columns=(INSTALLATION_CITY_COLUMN,),
    )
    return_device_counts = write_filtered_sheet(
        summary_wb=summary_wb,
        source_file=return_device_source_file,
        sheet_name=RETURN_DEVICE_SHEET_NAME,
        sheet_index=4,
        filter_column=OPERATION_TYPE_COLUMN,
        remove_blank=True,
    )
    removal_summary_counts = build_removal_summary_sheet(summary_wb)
    city_summary_count = build_city_recovery_summary_sheet(summary_wb)

    summary_wb.save(output_file)
    summary_wb.close()
    return (
        line_counts,
        service_counts,
        return_device_counts,
        removal_summary_counts,
        city_summary_count,
    )


# ============================================================
# 第二步：设备匹配、去重与结果汇总
# ============================================================

def normalize_order_number(value):
    """统一工单号格式，兼容Excel中的整数和文本编号。"""
    if is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def normalize_material_name(value):
    """清理物料名称前后空格，用于精确匹配。"""
    return "" if is_blank(value) else str(value).strip()

def normalize_text(value):
    """将筛选字段统一为去除前后空格的文本。"""
    return "" if is_blank(value) else str(value).strip()

def is_zero_value(value):
    """判断型号匹配是否为数值0或文本0。"""
    if isinstance(value, (int, float)):
        return value == 0
    return normalize_text(value) in {"0", "0.0"}

def is_matched_value(value):
    """判断匹配结果是否有效，空值和#N/A均视为未匹配。"""
    return normalize_text(value) not in {"", "#N/A"}

def ensure_output_columns(ws):
    """新增或复用结果列，避免脚本重复运行时产生同名列。"""
    existing = {
        normalize_header(cell.value): cell.column
        for cell in ws[1]
        if not is_blank(cell.value)
    }
    style_column = ws.max_column
    result = {}

    for header in NEW_COLUMNS:
        column = existing.get(header)
        if column is None:
            column = ws.max_column + 1
            copy_style(ws.cell(1, style_column), ws.cell(1, column))
            ws.cell(1, column).value = header
            ws.column_dimensions[ws.cell(1, column).column_letter].width = 22
        result[header] = column

    return result, style_column

def build_order_lookup(ws):
    """以工单号为键，建立业务类型查询表。"""
    columns = get_column_indexes(
        ws,
        (ORDER_NUMBER_COLUMN, BUSINESS_TYPE_1_COLUMN, BUSINESS_TYPE_2_COLUMN),
    )
    lookup = {}

    for row in range(2, ws.max_row + 1):
        order_number = ws.cell(row, columns[ORDER_NUMBER_COLUMN]).value
        key = normalize_order_number(order_number)
        if not key or key in lookup:
            continue
        lookup[key] = (
            order_number,
            ws.cell(row, columns[BUSINESS_TYPE_1_COLUMN]).value,
            ws.cell(row, columns[BUSINESS_TYPE_2_COLUMN]).value,
        )

    return lookup

def build_material_name_lookup(ws):
    """建立“物料名称 -> 原名称、划分类型”查询表。"""
    columns = get_column_indexes(ws, (MATERIAL_NAME_COLUMN, DIVISION_TYPE_COLUMN))
    lookup = {}

    for row in range(2, ws.max_row + 1):
        material_name = ws.cell(row, columns[MATERIAL_NAME_COLUMN]).value
        key = normalize_material_name(material_name)
        if not key or key in lookup:
            continue

        division_type = ws.cell(row, columns[DIVISION_TYPE_COLUMN]).value
        if key in ZERO_DIVISION_TYPE_MATERIALS:
            division_type = 0
        lookup[key] = (material_name, division_type)

    return lookup

def build_category_lookup(ws):
    """建立“物料名称 -> 分类”查询表。"""
    columns = get_column_indexes(ws, (MATERIAL_NAME_COLUMN, CATEGORY_COLUMN))
    lookup = {}

    for row in range(2, ws.max_row + 1):
        material_name = ws.cell(row, columns[MATERIAL_NAME_COLUMN]).value
        key = normalize_material_name(material_name)
        if not key or key in lookup:
            continue
        lookup[key] = ws.cell(row, columns[CATEGORY_COLUMN]).value

    return lookup

def unique_rows_by_order(ws, rows, order_column):
    """按关联单据号去重，保留筛选后源表中第一次出现的记录。"""
    seen = set()
    selected_rows = []

    for row in rows:
        key = normalize_order_number(ws.cell(row, order_column).value)
        if key in seen:
            continue
        seen.add(key)
        selected_rows.append(row)

    return selected_rows


def create_result_sheet(wb, source_ws, sheet_name, rows, extra_column=None):
    """按照源表结构创建结果sheet；数据行只复制值以提升速度。"""
    if sheet_name in wb.sheetnames:
        wb.remove(wb[sheet_name])
    target_ws = wb.create_sheet(sheet_name)
    source_max_column = source_ws.max_column

    copy_sheet_dimensions(source_ws, target_ws, copy_rows=False)

    for column in range(1, source_max_column + 1):
        copy_cell(source_ws.cell(1, column), target_ws.cell(1, column))

    if extra_column:
        extra_index = source_max_column + 1
        copy_style(source_ws.cell(1, source_max_column), target_ws.cell(1, extra_index))
        target_ws.cell(1, extra_index).value = extra_column
        target_ws.column_dimensions[
            target_ws.cell(1, extra_index).column_letter
        ].width = 22

    for source_row in rows:
        target_ws.append(
            [
                source_ws.cell(source_row, column).value
                for column in range(1, source_max_column + 1)
            ]
        )

    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
    last_column = source_max_column + (1 if extra_column else 0)
    target_ws.auto_filter.ref = (
        f"A1:{target_ws.cell(max(1, len(rows) + 1), last_column).coordinate}"
    )
    return target_ws

def append_rows(source_ws, target_ws, rows):
    """批量追加数据值，避免逐行计算max_row造成性能下降。"""
    source_max_column = source_ws.max_column
    for source_row in rows:
        target_ws.append(
            [
                source_ws.cell(source_row, column).value
                for column in range(1, source_max_column + 1)
            ]
        )

    target_ws.auto_filter.ref = (
        f"A1:{target_ws.cell(target_ws.max_row, source_max_column).coordinate}"
    )

def build_matched_return_sheets(wb):
    """严格按照sheet1、sheet2、sheet3的筛选顺序生成并合并数据。"""
    source_ws = wb[RETURN_DEVICE_SHEET]
    source_max_row = source_ws.max_row
    columns = get_column_indexes(
        source_ws,
        (
            RELATED_ORDER_COLUMN,
            ORDER_MATCH_COLUMN,
            BUSINESS_CATEGORY_1_COLUMN,
            MATERIAL_NAME_MATCH_COLUMN,
            MODEL_MATCH_COLUMN,
            CATEGORY_COLUMN,
        ),
    )
    all_rows = range(2, source_max_row + 1)

    # 一次遍历同时完成四组筛选，避免重复扫描整个工作表。
    sheet1_rows = []
    sheet2_zero_rows = []
    service_rows = []
    video_rows = []
    sheet3_zero_rows = []

    for row in all_rows:
        category = normalize_text(source_ws.cell(row, columns[CATEGORY_COLUMN]).value)
        material_match = source_ws.cell(row, columns[MATERIAL_NAME_MATCH_COLUMN]).value
        order_match = source_ws.cell(row, columns[ORDER_MATCH_COLUMN]).value
        business_type = normalize_text(
            source_ws.cell(row, columns[BUSINESS_CATEGORY_1_COLUMN]).value
        )
        model_match = source_ws.cell(row, columns[MODEL_MATCH_COLUMN]).value
        is_valid_row = (
            category == "终端"
            and is_matched_value(material_match)
            and is_matched_value(order_match)
        )

        if (
            is_valid_row
            and business_type not in BASE_EXCLUDED_BUSINESS_TYPES
            and not is_zero_value(model_match)
        ):
            sheet1_rows.append(row)

        if (
            is_valid_row
            and business_type not in BASE_EXCLUDED_BUSINESS_TYPES
            and is_zero_value(model_match)
        ):
            sheet2_zero_rows.append(row)

        model_text = normalize_text(model_match)
        if (
            is_valid_row
            and business_type in SERVICE_BUSINESS_TYPES
            and model_text in SERVICE_MODEL_TYPES
        ):
            service_rows.append(row)

        if is_valid_row and business_type == "视频监控":
            if is_zero_value(model_match):
                sheet3_zero_rows.append(row)
            else:
                video_rows.append(row)

        if (row - 1) % PROGRESS_INTERVAL == 0:
            print(f"结果表筛选进度：{row - 1}/{source_max_row - 1}", flush=True)

    print("正在生成匹配后拆回设备清单……", flush=True)
    sheet1_ws = create_result_sheet(
        wb, source_ws, MATCHED_RETURN_SHEET, sheet1_rows
    )

    # sheet2：型号匹配为0，按关联单据号去重，再剔除已存在于sheet1的关联单据号。
    deduplicated_sheet2_rows = unique_rows_by_order(
        source_ws, sheet2_zero_rows, columns[RELATED_ORDER_COLUMN]
    )
    sheet1_orders = {
        normalize_order_number(source_ws.cell(row, columns[RELATED_ORDER_COLUMN]).value)
        for row in sheet1_rows
    }
    sheet2_rows = [
        row
        for row in deduplicated_sheet2_rows
        if normalize_order_number(
            source_ws.cell(row, columns[RELATED_ORDER_COLUMN]).value
        )
        not in sheet1_orders
    ]
    create_result_sheet(
        wb,
        source_ws,
        SHEET2_NAME,
        sheet2_rows,
        extra_column=SHEET1_MATCH_COLUMN,
    )
    append_rows(source_ws, sheet1_ws, sheet2_rows)
    print(f"sheet2处理完成：{len(sheet2_rows)} 条", flush=True)

    # 重新读取初始清单：追加两类服务业务中指定型号的记录。
    append_rows(source_ws, sheet1_ws, service_rows)
    print(f"服务类数据追加完成：{len(service_rows)} 条", flush=True)

    # 重新读取初始清单：追加视频监控中型号匹配不为0的记录。
    append_rows(source_ws, sheet1_ws, video_rows)
    print(f"视频监控数据追加完成：{len(video_rows)} 条", flush=True)

    # sheet3：保留视频监控筛选，改筛型号匹配为0并按关联单据号去重。
    sheet3_rows = unique_rows_by_order(
        source_ws, sheet3_zero_rows, columns[RELATED_ORDER_COLUMN]
    )
    create_result_sheet(wb, source_ws, SHEET3_NAME, sheet3_rows)
    append_rows(source_ws, sheet1_ws, sheet3_rows)
    print(f"sheet3处理完成：{len(sheet3_rows)} 条", flush=True)

    return {
        "sheet1_initial": len(sheet1_rows),
        "sheet2": len(sheet2_rows),
        "service": len(service_rows),
        "video": len(video_rows),
        "sheet3": len(sheet3_rows),
        "final": sheet1_ws.max_row - 1,
    }

def fill_city_recovered_device_counts(wb):
    """按地市写入已拆回设备数量及终端回收率。"""
    if CITY_SUMMARY_SHEET not in wb.sheetnames:
        raise ValueError(f"工作簿中未找到工作表：{CITY_SUMMARY_SHEET}")

    source_ws = wb[MATCHED_RETURN_SHEET]
    source_columns = get_column_indexes(
        source_ws, (CITY_COLUMN, TOTAL_QUANTITY_COLUMN)
    )
    city_totals = {city: 0 for city in STANDARD_CITIES}

    for row in range(2, source_ws.max_row + 1):
        city = normalize_city(source_ws.cell(row, source_columns[CITY_COLUMN]).value)
        if city in city_totals:
            city_totals[city] += to_number(
                source_ws.cell(row, source_columns[TOTAL_QUANTITY_COLUMN]).value
            )

    summary_ws = wb[CITY_SUMMARY_SHEET]
    summary_columns = get_column_indexes(
        summary_ws,
        (
            CITY_COLUMN,
            EXPECTED_DEVICE_COUNT_COLUMN,
            RECOVERED_DEVICE_COUNT_COLUMN,
            RECOVERY_RATE_COLUMN,
        ),
    )
    expected_letter = get_column_letter(
        summary_columns[EXPECTED_DEVICE_COUNT_COLUMN]
    )
    recovered_letter = get_column_letter(
        summary_columns[RECOVERED_DEVICE_COUNT_COLUMN]
    )

    for row in range(2, summary_ws.max_row + 1):
        city = normalize_city(summary_ws.cell(row, summary_columns[CITY_COLUMN]).value)
        total = city_totals.get(city, 0)
        if isinstance(total, float) and total.is_integer():
            total = int(total)
        summary_ws.cell(
            row, summary_columns[RECOVERED_DEVICE_COUNT_COLUMN]
        ).value = total
        rate_cell = summary_ws.cell(row, summary_columns[RECOVERY_RATE_COLUMN])
        rate_cell.value = (
            f"=IFERROR({recovered_letter}{row}/{expected_letter}{row},0)"
        )
        rate_cell.number_format = "0.00%"

    return city_totals

def build_detail_summary_sheet(wb):
    """按业务类型2和地市汇总五类型号的总数量，并写入sheet1。"""
    source_ws = wb[MATCHED_RETURN_SHEET]
    columns = get_column_indexes(
        source_ws,
        (
            MODEL_MATCH_COLUMN,
            BUSINESS_CATEGORY_2_COLUMN,
            CITY_COLUMN,
            TOTAL_QUANTITY_COLUMN,
        ),
    )
    business_totals = defaultdict(float)
    city_totals = defaultdict(float)

    for row in range(2, source_ws.max_row + 1):
        model = normalize_text(source_ws.cell(row, columns[MODEL_MATCH_COLUMN]).value)
        if model not in SUMMARY_MODEL_TYPES:
            continue

        quantity = to_number(
            source_ws.cell(row, columns[TOTAL_QUANTITY_COLUMN]).value
        )
        business_type = normalize_text(
            source_ws.cell(row, columns[BUSINESS_CATEGORY_2_COLUMN]).value
        )
        city = normalize_city(source_ws.cell(row, columns[CITY_COLUMN]).value)

        if business_type in SUMMARY_BUSINESS_TYPES:
            business_totals[model, business_type] += quantity
        if city in STANDARD_CITIES:
            city_totals[model, city] += quantity

    if DETAIL_SUMMARY_SHEET in wb.sheetnames:
        wb.remove(wb[DETAIL_SUMMARY_SHEET])
    summary_ws = wb.create_sheet(DETAIL_SUMMARY_SHEET)

    def write_table(start_row, column_items, totals):
        """写入一张交叉汇总表，column_items为(显示名, 匹配值)。"""
        headers = ["型号匹配", *(label for label, _ in column_items), "总计"]
        for column, header in enumerate(headers, 1):
            summary_ws.cell(start_row, column).value = header

        for row_offset, model in enumerate(SUMMARY_MODEL_TYPES, 1):
            target_row = start_row + row_offset
            summary_ws.cell(target_row, 1).value = model
            row_total = 0
            for column, (_, key) in enumerate(column_items, 2):
                value = totals[model, key]
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                summary_ws.cell(target_row, column).value = value
                row_total += value
            summary_ws.cell(target_row, len(headers)).value = row_total

        total_row = start_row + len(SUMMARY_MODEL_TYPES) + 1
        summary_ws.cell(total_row, 1).value = "总计"
        grand_total = 0
        for column, (_, key) in enumerate(column_items, 2):
            value = sum(totals[model, key] for model in SUMMARY_MODEL_TYPES)
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            summary_ws.cell(total_row, column).value = value
            grand_total += value
        summary_ws.cell(total_row, len(headers)).value = grand_total
        return total_row

    business_columns = tuple((name, name) for name in SUMMARY_BUSINESS_TYPES)
    first_table_end = write_table(1, business_columns, business_totals)
    second_table_start = first_table_end + 3
    second_table_end = write_table(second_table_start, SUMMARY_CITIES, city_totals)

    summary_ws.freeze_panes = "B2"
    summary_ws.column_dimensions["A"].width = 16
    for column in range(2, summary_ws.max_column + 1):
        summary_ws.column_dimensions[get_column_letter(column)].width = 16

    # sheet1为独立新建的汇总表，不继承其他sheet的百分比等数字格式。
    for row in summary_ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0"

    return {
        "business_grand_total": sum(business_totals.values()),
        "city_grand_total": sum(city_totals.values()),
    }

def update_return_device_sheet(workbook_file, output_file=None):
    """匹配拆除工单和物料信息，并写入一体化拆回设备清单。"""
    if output_file is None:
        output_file = workbook_file
    print(f"正在读取工作簿：{workbook_file}", flush=True)
    wb = load_workbook(workbook_file)
    try:
        missing_sheets = [
            name
            for name in (
                RETURN_DEVICE_SHEET,
                REMOVAL_SUMMARY_SHEET,
                MATERIAL_NAME_SHEET,
                PROVINCE_MATERIAL_SHEET,
                CITY_SUMMARY_SHEET,
            )
            if name not in wb.sheetnames
        ]
        if missing_sheets:
            raise ValueError(f"工作簿中未找到工作表：{', '.join(missing_sheets)}")

        return_ws = wb[RETURN_DEVICE_SHEET]
        removal_ws = wb[REMOVAL_SUMMARY_SHEET]
        material_name_ws = wb[MATERIAL_NAME_SHEET]
        province_material_ws = wb[PROVINCE_MATERIAL_SHEET]
        return_columns = get_column_indexes(
            return_ws, (RELATED_ORDER_COLUMN, MATERIAL_NAME_COLUMN)
        )
        output_columns, _ = ensure_output_columns(return_ws)
        print("正在建立工单和物料匹配索引……", flush=True)
        order_lookup = build_order_lookup(removal_ws)
        material_name_lookup = build_material_name_lookup(material_name_ws)
        category_lookup = build_category_lookup(province_material_ws)

        order_matched_count = 0
        order_unmatched_count = 0
        material_matched_count = 0
        category_matched_count = 0

        total_rows = return_ws.max_row - 1
        print(f"开始匹配一体化拆回设备清单，共 {total_rows} 条。", flush=True)
        for row in range(2, return_ws.max_row + 1):
            for column in output_columns.values():
                return_ws.cell(row, column).value = None

            related_order = return_ws.cell(
                row, return_columns[RELATED_ORDER_COLUMN]
            ).value
            match = order_lookup.get(normalize_order_number(related_order))
            if match is None:
                order_unmatched_count += 1
            else:
                return_ws.cell(row, output_columns[ORDER_MATCH_COLUMN]).value = match[0]
                return_ws.cell(
                    row, output_columns[BUSINESS_CATEGORY_1_COLUMN]
                ).value = match[1]
                return_ws.cell(
                    row, output_columns[BUSINESS_CATEGORY_2_COLUMN]
                ).value = match[2]
                order_matched_count += 1

            material_name = return_ws.cell(
                row, return_columns[MATERIAL_NAME_COLUMN]
            ).value
            material_key = normalize_material_name(material_name)

            material_match = material_name_lookup.get(material_key)
            if material_match is not None:
                return_ws.cell(
                    row, output_columns[MATERIAL_NAME_MATCH_COLUMN]
                ).value = material_match[0]
                return_ws.cell(
                    row, output_columns[MODEL_MATCH_COLUMN]
                ).value = material_match[1]
                material_matched_count += 1

            if material_key in category_lookup:
                return_ws.cell(
                    row, output_columns[CATEGORY_COLUMN]
                ).value = category_lookup[material_key]
                category_matched_count += 1

            if (row - 1) % PROGRESS_INTERVAL == 0:
                print(f"匹配进度：{row - 1}/{total_rows}", flush=True)

        print("基础匹配完成，开始生成结果sheet。", flush=True)
        result_sheet_counts = build_matched_return_sheets(wb)
        print("正在生成sheet1的业务类型及地市汇总表……", flush=True)
        detail_summary_counts = build_detail_summary_sheet(wb)
        print("正在按地市汇总已拆回设备数量……", flush=True)
        city_recovered_counts = fill_city_recovered_device_counts(wb)
        print(f"正在保存工作簿：{output_file}", flush=True)
        wb.save(output_file)
        print("工作簿保存完成。", flush=True)
    finally:
        wb.close()

    return {
        "order_matched": order_matched_count,
        "order_unmatched": order_unmatched_count,
        "material_matched": material_matched_count,
        "category_matched": category_matched_count,
        "result_sheets": result_sheet_counts,
        "detail_summary_counts": detail_summary_counts,
        "city_recovered_counts": city_recovered_counts,
    }


# ============================================================
# 统一执行入口
# ============================================================
def parse_args():
    parser = argparse.ArgumentParser(description="生成指定月份的终端回收率汇总表")
    parser.add_argument("--start-date", required=True)
    parser.add_argument("--end-date", required=True)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--district-map", type=Path, help="含地市、区县、标准区县的映射表；默认在输出目录查找")
    parser.add_argument("--mode", choices=("file", "database", "both"), default="both")
    parser.add_argument("--database", type=Path, help="兼容旧命令；始终使用 database.py 中的 MySQL 配置")
    return parser.parse_args()


def validate_dates(start_date, end_date):
    from datetime import date
    start, end = date.fromisoformat(start_date), date.fromisoformat(end_date)
    if start.isoformat() != start_date or end.isoformat() != end_date:
        raise ValueError("Dates must use YYYY-MM-DD")
    if start > end or start.strftime("%Y-%m") != end.strftime("%Y-%m"):
        raise ValueError("Date range must be ordered and within one month")
    return start.strftime("%Y-%m")


def export(start_date, end_date, input_dir=DEFAULT_INPUT_DIR, output_dir=DEFAULT_OUTPUT_DIR):
    month = validate_dates(start_date, end_date)
    input_dir, output_dir = Path(input_dir), Path(output_dir)
    sources = [input_dir / path for path in get_business_source_files(month)]
    material = input_dir / "物料名称表.xlsx"
    if not material.is_file():
        material = MATERIAL_NAME_SOURCE_FILE
    baseline = input_dir / "全省物资基准库.xlsx"
    for source in [*sources, material, baseline]:
        if not source.is_file():
            raise FileNotFoundError(source)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / get_output_file(month)
    build_workbook(month, *sources, material_name_source_file=material,
                   province_material_source_file=baseline, output_file=output_file)
    stats = update_return_device_sheet(output_file, output_file)
    return output_file, stats

def export_database(start_date, end_date, database, output_dir=DEFAULT_OUTPUT_DIR, mode="both", district_map=None):
    import json
    import sys
    import tempfile
    sys.path.insert(0, str(SCRIPT_DIR.parents[1]))
    from storage.terminal_recovery import restore_sources, source_run_ids
    from metrics.terminal_recovery.report import read_report
    from metrics.installation.common import save_metric
    from metrics.terminal_recovery.districts import resolve_mapping, load_mapping, build_district_report, metric_rows

    validate_dates(start_date, end_date)
    mapping_file = resolve_mapping(output_dir, district_map)
    load_mapping(mapping_file)
    if mode not in {"file", "database", "both"}:
        raise ValueError("Unsupported mode")
    # Temporary Excel files adapt database snapshots to the unchanged legacy engine.
    with tempfile.TemporaryDirectory(prefix="terminal_recovery_") as directory:
        print("正在从数据库恢复源快照……", flush=True)
        source_dir = Path(directory) / "sources"
        snapshots = restore_sources(database, start_date, end_date, source_dir)
        print("源快照恢复完成，正在运行原汇总算法……", flush=True)
        target_dir = output_dir if mode in {"file", "both"} else Path(directory) / "report"
        output_file, stats = export(start_date, end_date, source_dir, target_dir)
        source_runs = source_run_ids(database, snapshots)
        report = read_report(output_file, start_date, end_date, snapshots, source_runs)
        print("正在清洗区县并计算区县回收率……", flush=True)
        district_file = Path(target_dir) / f"终端回收区县结果_{start_date[:7]}.xlsx"
        district_args = (
            source_dir / f"一体化专线拆机清单_{start_date[:7]}.xlsx",
            source_dir / f"服务类产品支撑工单_{start_date[:7]}.xlsx",
            output_file, mapping_file,
        )
        try:
            report["districts"] = build_district_report(*district_args, district_file)
        except PermissionError:
            district_file = Path(target_dir) / f"终端回收区县结果_口径修正_{start_date[:7]}.xlsx"
            print(f"原区县结果表正在使用，改为输出：{district_file}", flush=True)
            report["districts"] = build_district_report(*district_args, district_file)
        report["metric_version"] = "3.0"
        report["results"].extend(metric_rows(report["districts"]))
        stats["district_file"] = str(district_file.resolve()) if mode in {"file", "both"} else None
        report["metric_run_id"] = None
        if mode in {"database", "both"}:
            print("正在将地市、设备和区县指标写入数据库……", flush=True)
            report["metric_run_id"] = save_metric(database, report, source_runs)
        stats["metric_run_id"] = report["metric_run_id"]
        stats["result_count"] = len(report["results"])
        stats["json_file"] = None
        if mode in {"file", "both"}:
            json_file = output_file.with_suffix(".json")
            json_file.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
            stats["json_file"] = str(json_file.resolve())
    return (output_file if mode in {"file", "both"} else None), stats


def run_export(start_date, end_date, *, mode="both", database=None,
               output_dir=DEFAULT_OUTPUT_DIR, district_map=None):
    return export_database(start_date, end_date, database, output_dir, mode, district_map)


def main():
    args = parse_args()
    output_file, stats = run_export(**vars(args))
    if output_file is not None:
        print(f"全部处理完成，输出文件：{output_file.resolve()}")
    else:
        from storage.database import _database_name
        print(f"全部处理完成，两个汇总 sheet 的指标已存入 MySQL：{_database_name()}")
    print(stats)


if __name__ == "__main__":
    main()
