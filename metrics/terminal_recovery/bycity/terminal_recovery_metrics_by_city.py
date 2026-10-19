#!/usr/bin/env python3
"""从省级终端回收成功批次提取指定地市及其标准区县指标。"""

from __future__ import annotations

import argparse
import calendar
import json
import math
import re
import sys
import tempfile
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from storage.database import connect, initialize
from storage.metric_results import insert_results
from storage.terminal_recovery import restore_sources


METRIC_CODE = "terminal_recovery_metrics_by_city"
SOURCE_METRIC_CODE = "terminal_recovery"
PREPARATION_METRIC_CODE = "terminal_recovery_bycity_preparation"
RESULT_OWNER = "terminal_recovery"
METRIC_VERSION = "1.2.0"
MATCHED_SHEET = "匹配后拆回设备清单（三类标签已去重）"
DETAIL_SHEET = "sheet1"


def normalize_city(value: object) -> str:
    city = "" if value is None else str(value).strip()
    return city.removesuffix("市")


def display_city(value: object) -> str:
    return f"{normalize_city(value)}市"


def month_bounds(month: str) -> tuple[str, str]:
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise ValueError("月份格式应为 YYYY-MM")
    year, number = map(int, month.split("-"))
    if not 1 <= number <= 12:
        raise ValueError("月份格式应为 YYYY-MM")
    return f"{month}-01", f"{month}-{calendar.monthrange(year, number)[1]:02d}"


def _parse_row(row: dict) -> dict:
    dimension = json.loads(row["dimension_value"])
    return {
        "metric_code": row["metric_code"],
        "dimension_type": row["dimension_type"],
        "dimension": dimension,
        "numerator": row["numerator"],
        "denominator": row["denominator"],
        "metric_value": row["metric_value"],
    }


def _belongs_to_city(item: dict, city: str) -> bool:
    """保留省级对比维度，以及目标地市的区县和业务类型明细。"""
    kind = item["dimension_type"]
    if kind in {"city_summary", "city"}:
        return True
    if kind == "business":
        return True
    label = str(item["dimension"].get("label") or "")
    if kind == "district_summary":
        return label.startswith(city + "·") or label.startswith(display_city(city) + "·")
    return False


def _build_all_city_business(database: Path | None, start: str, end: str) -> tuple[list[dict], dict]:
    """从数据库原始快照重建一次全省匹配清单，生成所有地市业务设备维度。"""
    from openpyxl import load_workbook
    from metrics.terminal_recovery.report import table
    from metrics.terminal_recovery.terminal_recovery_export_online import export

    with tempfile.TemporaryDirectory(prefix="terminal_bycity_") as directory:
        root = Path(directory)
        source_dir = root / "sources"
        output_dir = root / "report"
        snapshots = restore_sources(database, start, end, source_dir)
        output_file, _stats = export(start, end, source_dir, output_dir)
        workbook = load_workbook(output_file, read_only=True, data_only=False)
        try:
            detail_rows = list(workbook[DETAIL_SHEET].iter_rows(values_only=True))
            headers = [index for index, row in enumerate(detail_rows) if row[0] == "型号匹配"]
            if len(headers) != 2:
                raise ValueError("sheet1 must contain exactly two cross tables")
            business_table = table(detail_rows, headers[0], headers[1])
            business_columns = [column for column in business_table["columns"][1:] if column != "总计"]
            device_labels = [
                row[business_table["columns"][0]] for row in business_table["rows"]
                if row[business_table["columns"][0]] != "总计"
            ]

            matched_rows = workbook[MATCHED_SHEET].iter_rows(values_only=True)
            matched_headers = list(next(matched_rows))
            required = {"地市", "业务类别2", "型号匹配", "总数量"}
            if not required.issubset(matched_headers):
                raise ValueError(f"匹配后拆回设备清单缺少字段：{sorted(required - set(matched_headers))}")
            positions = {name: matched_headers.index(name) for name in required}
            counts = Counter()
            municipalities = []
            matched_source_rows = 0
            for values in matched_rows:
                municipality = normalize_city(values[positions["地市"]])
                business_type = str(values[positions["业务类别2"]] or "").strip()
                device = str(values[positions["型号匹配"]] or "").strip()
                quantity = values[positions["总数量"]] or 0
                if municipality and business_type and device:
                    if not isinstance(quantity, (int, float)) or not math.isfinite(quantity):
                        raise ValueError(f"总数量不是有效数值：{quantity!r}")
                    if municipality not in municipalities:
                        municipalities.append(municipality)
                    counts[municipality, device, business_type] += quantity
                    matched_source_rows += 1
        finally:
            workbook.close()

    results = []
    for municipality in municipalities:
        for row_index, device in enumerate(device_labels, 1):
            for column_index, business_type in enumerate(business_columns, 2):
                value = counts[municipality, device, business_type]
                results.append({
                    "metric_code": "terminal_recovery_count",
                    "dimension_type": "city_business",
                    "dimension": {
                        "sheet": MATCHED_SHEET, "table": "city_business",
                        "row_index": row_index, "column_index": column_index,
                        "label_column": "地市·型号匹配",
                        "label": f"{municipality}·{device}", "column": business_type,
                    },
                    "numerator": value, "denominator": 1, "metric_value": value,
                })
    return results, {
        "source_snapshots": snapshots,
        "matched_source_rows": matched_source_rows,
        "city_count": len(municipalities),
        "business_type_count": len(business_columns),
        "device_type_count": len(device_labels),
    }


def _save_preparation(database: Path | None, start: str, end: str,
                      source_run_id: str, results: list[dict], quality: dict) -> str:
    run_id = f"metric_terminal_city_prep_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}"
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    source_runs = [source_run_id, "scope:all_cities", json.dumps(quality, ensure_ascii=False)]
    with connect(database) as connection:
        connection.execute(
            """INSERT INTO metric_run(metric_run_id,metric_code,metric_version,
               period_start,period_end,started_at,status,source_runs)
               VALUES(?,?,?,?,?,?,'running',?)""",
            (run_id, PREPARATION_METRIC_CODE, METRIC_VERSION, start, end, now,
             json.dumps(source_runs, ensure_ascii=False)),
        )
    try:
        with connect(database) as connection:
            insert_results(connection, RESULT_OWNER, run_id, results)
            connection.execute(
                "UPDATE metric_run SET finished_at=?,status='success' WHERE metric_run_id=?",
                (datetime.now(timezone.utc).isoformat(timespec="seconds"), run_id),
            )
    except Exception as exc:
        with connect(database) as connection:
            connection.execute(
                "UPDATE metric_run SET finished_at=?,status='failed',error_message=? WHERE metric_run_id=?",
                (datetime.now(timezone.utc).isoformat(timespec="seconds"), str(exc), run_id),
            )
        raise
    return run_id


def _load_or_prepare_city_business(database: Path | None, start: str, end: str,
                                   source_run_id: str) -> tuple[list[dict], dict]:
    with connect(database) as connection:
        runs = connection.execute(
            """SELECT * FROM metric_run WHERE metric_code=? AND status='success'
               AND period_start=? AND period_end=?
               ORDER BY started_at DESC,metric_run_id DESC""",
            (PREPARATION_METRIC_CODE, start, end),
        ).fetchall()
        for run in runs:
            try:
                sources = json.loads(run.get("source_runs") or "[]")
            except (TypeError, json.JSONDecodeError):
                sources = []
            if source_run_id not in sources:
                continue
            rows = connection.execute(
                """SELECT metric_code,dimension_type,dimension_value,
                          numerator,denominator,metric_value
                   FROM result_terminal_recovery WHERE metric_run_id=? ORDER BY result_id""",
                (run["metric_run_id"],),
            ).fetchall()
            results = [_parse_row(dict(row)) for row in rows]
            return results, {"preparation_metric_run_id": run["metric_run_id"], "cache_hit": True}

    results, quality = _build_all_city_business(database, start, end)
    preparation_run_id = _save_preparation(database, start, end, source_run_id, results, quality)
    return results, {
        "preparation_metric_run_id": preparation_run_id,
        "cache_hit": False,
        **quality,
    }

def calculate(database: Path | None, month: str, city: str) -> dict:
    city_short = normalize_city(city)
    if not city_short:
        raise ValueError("--city 不能为空")
    period_start, period_end = month_bounds(month)
    with connect(database) as connection:
        source_run = connection.execute(
            """SELECT * FROM metric_run
               WHERE metric_code=? AND status='success'
                 AND period_start=? AND period_end=?
               ORDER BY started_at DESC, metric_run_id DESC LIMIT 1""",
            (SOURCE_METRIC_CODE, period_start, period_end),
        ).fetchone()
        if source_run is None:
            raise RuntimeError(f"没有找到 {month} 的成功省级终端回收指标批次")
        rows = connection.execute(
            """SELECT metric_code,dimension_type,dimension_value,
                      numerator,denominator,metric_value
               FROM result_terminal_recovery
               WHERE metric_run_id=? ORDER BY result_id""",
            (source_run["metric_run_id"],),
        ).fetchall()

    source_results = [_parse_row(dict(row)) for row in rows]
    results = [item for item in source_results if _belongs_to_city(item, city_short)]
    all_city_business, city_business_quality = _load_or_prepare_city_business(
        database, period_start, period_end, source_run["metric_run_id"]
    )
    city_business = [
        item for item in all_city_business
        if str(item["dimension"].get("label") or "").startswith(city_short + "·")
    ]
    results.extend(city_business)
    city_markers = [
        item for item in results
        if item["dimension_type"] in {"district_summary", "city_business"}
    ]
    if not city_markers:
        raise RuntimeError(f"省级终端回收批次中没有 {display_city(city_short)} 的区县或业务明细")
    for item in results:
        for field in ("numerator", "denominator", "metric_value"):
            value = item[field]
            if value is not None and (not isinstance(value, (int, float)) or not math.isfinite(value)):
                raise ValueError(f"终端回收结果包含无效数值：{field}={value!r}")
    return {
        "metric_code": METRIC_CODE,
        "metric_version": METRIC_VERSION,
        "month": month,
        "city": display_city(city_short),
        "period_start": period_start,
        "period_end": period_end,
        "source_metric_run_id": source_run["metric_run_id"],
        "quality": {
            "source_result_count": len(source_results),
            "selected_result_count": len(results),
            "city_detail_result_count": len(city_markers),
            "city_business": city_business_quality,
        },
        "results": results,
    }


def save_results(database: Path | None, report: dict) -> str:
    city_short = normalize_city(report["city"])
    run_id = f"metric_terminal_city_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}"
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    source_runs = [report["source_metric_run_id"], f"city:{city_short}"]
    with connect(database) as connection:
        connection.execute(
            """INSERT INTO metric_run(metric_run_id,metric_code,metric_version,
               period_start,period_end,started_at,status,source_runs)
               VALUES(?,?,?,?,?,?,'running',?)""",
            (run_id, METRIC_CODE, METRIC_VERSION, report["period_start"], report["period_end"],
             now, json.dumps(source_runs, ensure_ascii=False)),
        )
    try:
        with connect(database) as connection:
            insert_results(connection, RESULT_OWNER, run_id, report["results"])
            connection.execute(
                "UPDATE metric_run SET finished_at=?,status='success' WHERE metric_run_id=?",
                (datetime.now(timezone.utc).isoformat(timespec="seconds"), run_id),
            )
    except Exception as exc:
        with connect(database) as connection:
            connection.execute(
                "UPDATE metric_run SET finished_at=?,status='failed',error_message=? WHERE metric_run_id=?",
                (datetime.now(timezone.utc).isoformat(timespec="seconds"), str(exc), run_id),
            )
        raise
    return run_id


def run(database: Path | None, month: str, city: str, mode: str = "both",
        output: Path | None = None) -> dict:
    initialize(database)
    report = calculate(database, month, city)
    report["metric_run_id"] = save_results(database, report) if mode in {"database", "both"} else None
    if mode in {"file", "both"}:
        target = output or Path(
            f"outputs/终端回收地市指标_{normalize_city(city)}_{month}.json"
        )
        target = target.expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        report["output_file"] = str(target)
    print(
        f"完成 {display_city(city)} {month} 终端回收指标："
        f"{report['metric_run_id'] or '仅文件'}",
        flush=True,
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="提取指定地市及标准区县的终端回收指标")
    parser.add_argument("--month", required=True, help="月报月份，格式 YYYY-MM")
    parser.add_argument("--city", required=True, help="目标地市，例如 杭州 或 杭州市")
    parser.add_argument("--mode", choices=("file", "database", "both"), default="both")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--database", type=Path)
    args = parser.parse_args()
    run(args.database, args.month, args.city, args.mode, args.output)


if __name__ == "__main__":
    main()

