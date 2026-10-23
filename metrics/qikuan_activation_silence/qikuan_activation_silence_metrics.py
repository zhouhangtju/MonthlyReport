#!/usr/bin/env python3
"""Calculate report-month broadband silence rates with T-3 orders as denominator."""

from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import json
import re
import sys
import uuid
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from storage.database import connect
from storage.metric_results import insert_results

OWNER = "qikuan_activation_silence"
CODE = "qikuan_activation_silence_rate"
SCHEMA = ROOT / "sql" / "create_qikuan_activation_silence.sql"
CITY_ORDER = ["杭州", "宁波", "温州", "嘉兴", "湖州", "绍兴", "金华", "衢州", "舟山", "台州", "丽水"]


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def bounds(month: str) -> tuple[str, str]:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise ValueError("--month 应为 YYYY-MM")
    year, number = map(int, month.split("-"))
    return f"{month}-01", f"{month}-{calendar.monthrange(year, number)[1]:02d}"


def cohort_month(report_month: str) -> str:
    """The October silence indicator uses broadband orders completed in July."""
    bounds(report_month)
    year, number = map(int, report_month.split("-"))
    year, zero_based = divmod(year * 12 + number - 1 - 3, 12)
    return f"{year:04d}-{zero_based + 1:02d}"


def order_id(value: object) -> str:
    if value is None:
        return ""
    value = str(value).strip()
    return value[:-2] if value.endswith(".0") and value[:-2].isdigit() else value


def clean(value: object) -> str:
    return str(value).strip() if value is not None else ""


def make_schema() -> None:
    with connect() as connection:
        connection.executescript(SCHEMA.read_text(encoding="utf-8"))


def import_manual(month: str, path: Path) -> dict:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    with connect() as connection:
        existing = connection.execute(
            "SELECT batch_id,row_count,duplicate_count FROM raw_qikuan_activation_silence_batch "
            "WHERE report_month=? AND file_sha256=?", (month, digest)
        ).fetchone()
    if existing:
        return dict(existing, reused=True)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(iterator)]
        required = ("工单号", "地市", "区县", "定单报结时间")
        if len(headers) != len(set(headers)) or any(name not in headers for name in required):
            raise ValueError(f"手工表缺少必要字段或有重复字段：{required}")
        index = {name: headers.index(name) for name in required}
        rows: dict[str, tuple] = {}
        duplicate_count = 0
        for source_row, values in enumerate(iterator, 2):
            if not any(value is not None for value in values):
                continue
            key = order_id(values[index["工单号"]])
            if not key:
                raise ValueError(f"手工表第 {source_row} 行工单号为空")
            completed = clean(values[index["定单报结时间"]])
            if completed[:7] != cohort_month(month):
                raise ValueError(
                    f"手工表第 {source_row} 行报结月份为 {completed[:7]}，"
                    f"与报告月 {month} 的 T-3 月份 {cohort_month(month)} 不符"
                )
            if key in rows:
                duplicate_count += 1
                continue  # 同一工单只计一次，保留源表第一条。
            payload = json.dumps(dict(zip(headers, values)), ensure_ascii=False, default=str)
            rows[key] = (key, source_row, clean(values[index["地市"]]),
                         clean(values[index["区县"]]), completed, payload)
    finally:
        workbook.close()
    if not rows:
        raise ValueError("手工表没有有效工单")
    batch_id = "silence_src_" + uuid.uuid4().hex
    with connect() as connection:
        connection.execute(
            "INSERT INTO raw_qikuan_activation_silence_batch "
            "(batch_id,report_month,source_file,file_sha256,imported_at,row_count,duplicate_count) "
            "VALUES (?,?,?,?,?,?,?)",
            (batch_id, month, str(path), digest, now(), len(rows), duplicate_count),
        )
        for row in rows.values():
            connection.execute(
                "INSERT INTO raw_qikuan_activation_silence "
                "(batch_id,order_id,source_row,manual_city,manual_district,completion_time,source_data) "
                "VALUES (?,?,?,?,?,?,?)", (batch_id, *row)
            )
    return {"batch_id": batch_id, "row_count": len(rows),
            "duplicate_count": duplicate_count, "reused": False}


def latest_sources(month: str) -> tuple[dict, dict]:
    start, end = bounds(cohort_month(month))
    with connect() as connection:
        online = connection.execute(
            "SELECT * FROM collection_run WHERE dataset_code='youshu_install_list' "
            "AND period_start=? AND period_end=? AND status='success' "
            "ORDER BY finished_at DESC,collection_run_id DESC LIMIT 1", (start, end)
        ).fetchone()
        manual = connection.execute(
            "SELECT * FROM raw_qikuan_activation_silence_batch WHERE report_month=? "
            "ORDER BY imported_at DESC,batch_id DESC LIMIT 1", (month,)
        ).fetchone()
    if not online or not manual:
        raise ValueError(f"{month} 缺少成功的企宽清单采集批次或手工沉默表入库批次")
    return dict(online), dict(manual)


def calculate(month: str) -> dict:
    online, manual = latest_sources(month)
    with connect() as connection:
        manual_rows = connection.execute(
            "SELECT order_id,manual_city,manual_district FROM raw_qikuan_activation_silence "
            "WHERE batch_id=?", (manual["batch_id"],)
        ).fetchall()
        source_rows = connection.execute(
            "SELECT `工单id` AS order_id,`地市` AS city,`区县` AS district,"
            "`市场网格` AS grid FROM youshu_install_list "
            "WHERE LEFT(`报结时间`,7)=?", (cohort_month(month),)
        ).fetchall()
    if not source_rows:
        raise ValueError(f"{month} 指标的 T-3 企宽装机清单分母为空")
    silent = {order_id(row["order_id"]): row for row in manual_rows}
    matched: set[str] = set()
    groups: dict[tuple, list[int]] = defaultdict(lambda: [0, 0])
    city_mismatch = 0
    district_mismatch = 0
    no_grid = 0
    for item in source_rows:
        key = order_id(item["order_id"])
        city = clean(item["city"]).removesuffix("市") or "（空）"
        district = clean(item["district"]) or "（空）"
        grid = clean(item["grid"])
        is_silent = int(key in silent)
        if is_silent:
            matched.add(key)
            manual_row = silent[key]
            city_mismatch += int(clean(manual_row["manual_city"]).removesuffix("市") != city)
            district_mismatch += int(clean(manual_row["manual_district"]) != district)
        dimensions = [("province",), ("city", city), ("district", city, district)]
        if grid:
            dimensions.append(("grid", city, grid))
        else:
            no_grid += 1
        for dimension in dimensions:
            groups[dimension][0] += is_silent
            groups[dimension][1] += 1

    results = []
    for dimension, (numerator, denominator) in groups.items():
        kind = dimension[0]
        dims = {"month": month}
        if kind in ("city", "district", "grid"):
            dims["city"] = dimension[1]
        if kind == "district":
            dims["district"] = dimension[2]
        if kind == "grid":
            dims["grid"] = dimension[2]
        results.append({"metric_code": CODE, "dimension_type": kind, "dimension": dims,
                        "numerator": numerator, "denominator": denominator,
                        "metric_value": numerator / denominator if denominator else None})
    unmatched = sorted(set(silent) - matched)
    audit = {"online_collection_run_id": online["collection_run_id"],
             "report_month": month, "denominator_month": cohort_month(month),
             "manual_batch_id": manual["batch_id"], "manual_file_sha256": manual["file_sha256"],
             "manual_unique_orders": len(silent), "matched_orders": len(matched),
             "unmatched_orders": len(unmatched), "online_missing_grid": no_grid,
             "matched_city_mismatches": city_mismatch,
             "matched_district_mismatches": district_mismatch,
             "denominator_rule": "T-3 complete broadband list; all action types",
             "grid_grouping": "city + market_grid",
             "grid_ranking_min_denominator": 5}
    return {"month": month, "results": results, "audit": audit, "unmatched": unmatched}


def save(report: dict) -> str:
    month = report["month"]
    start, end = bounds(month)
    run_id = "qikuan_silence_" + uuid.uuid4().hex
    with connect() as connection:
        connection.execute(
            "INSERT INTO metric_run (metric_run_id,metric_code,metric_version,period_start,period_end,"
            "started_at,status,source_runs) VALUES (?,?,?,?,?,?,'running',?)",
            (run_id, OWNER, "1.0.1", start, end, now(),
             json.dumps(report["audit"], ensure_ascii=False)),
        )
    try:
        with connect() as connection:
            insert_results(connection, OWNER, run_id, report["results"])
            connection.execute(
                "UPDATE metric_run SET status='success',finished_at=? WHERE metric_run_id=?",
                (now(), run_id),
            )
    except Exception as exc:
        with connect() as connection:
            connection.execute(
                "UPDATE metric_run SET status='failed',finished_at=?,error_message=? "
                "WHERE metric_run_id=?", (now(), str(exc), run_id)
            )
        raise
    return run_id


def run(month: str, input_file: Path | None, output_dir: Path) -> dict:
    make_schema()
    imported = import_manual(month, input_file) if input_file else None
    report = calculate(month)
    report["metric_run_id"] = save(report)
    output_dir = output_dir.expanduser().resolve() / month
    output_dir.mkdir(parents=True, exist_ok=True)
    unmatched_path = output_dir / "未匹配工单号.csv"
    with unmatched_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["工单号"])
        writer.writerows([value] for value in report.pop("unmatched"))
    output_path = output_dir / "企宽激活即沉默指标.json"
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    province = next(row for row in report["results"] if row["dimension_type"] == "province")
    return {"month": month, "manual_import": imported, "metric_run_id": report["metric_run_id"],
            "numerator": province["numerator"], "denominator": province["denominator"],
            "rate": province["metric_value"], "result_rows": len(report["results"]),
            "unmatched_orders": report["audit"]["unmatched_orders"],
            "output": str(output_path), "unmatched_output": str(unmatched_path)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--month", required=True, help="指标报告月份 YYYY-MM；分母自动取 T-3 自然月")
    parser.add_argument("--input", type=Path, help="手工上传的激活即沉默 Excel；不传则复用本月最新已入库批次")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs" / "qikuan_activation_silence")
    args = parser.parse_args()
    bounds(args.month)
    print(json.dumps(run(args.month, args.input, args.output_dir), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
