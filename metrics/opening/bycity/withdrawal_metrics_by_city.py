#!/usr/bin/env python3
"""Calculate the withdrawal page for one city and its standard counties."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from config.report_periods import integration_opening_period
from metrics.opening.withdrawal import calculate as calculate_line, load_rows, parse_time, text
from metrics.opening.bycity.dedicated_line_metrics_by_city import (
    DEFAULT_COUNTY_MAP, display_city, load_county_mapping, normalize_city,
    standardize_county, UNASSIGNED_COUNTY,
)
from storage.database import connect, initialize
from storage.metric_results import insert_results

OWNER = "business_support_withdrawal_by_city"
VERSION = "1.0.0"
PROJECT = Path(__file__).resolve().parents[3]
DEFAULT_LINE_REASONS = PROJECT / "data" / "offline_data" / "7.26-8.25撤退单原因需反馈.xlsx"
DEFAULT_BROADBAND = PROJECT / "data" / "offline_data" / "企宽.xlsx"
LINE_REASONS = {"网络建设原因": "network", "用户原因": "customer",
                "前台原因": "frontDesk", "其他原因": "other"}
BROADBAND_REASONS = {"用户原因": "customer", "前台原因": "frontDesk",
                     "建设原因": "construction", "其他原因": "other", "网络原因": "network"}


def read_sheet(path: Path, sheet: str, required: set[str]):
    if not path.is_file():
        raise FileNotFoundError(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book:
            raise ValueError(f"{path.name} 缺少工作表 {sheet}")
        rows = book[sheet].iter_rows(values_only=True)
        headers = [text(v) for v in next(rows)]
        missing = required - set(headers)
        if missing:
            raise ValueError(f"{path.name} 缺少列：{sorted(missing)}")
        positions = {key: headers.index(key) for key in required}
        for line, values in enumerate(rows, 2):
            if any(v is not None for v in values):
                yield line, {key: values[index] if index < len(values) else None
                             for key, index in positions.items()}
    finally:
        book.close()


def locate_cell(value: object, aliases_by_city: dict[str, dict[str, str]]):
    """Use CELL only; never use the source CITY as a fallback."""
    raw = text(value)
    if not raw:
        return None, None, "blank"
    matches = {(city, standard) for city, aliases in aliases_by_city.items()
               for alias, standard in aliases.items() if alias and alias in raw}
    if not matches:
        return None, None, "unmatched"
    if len(matches) != 1:
        return None, None, "conflict"
    city, county = next(iter(matches))
    return city, county, "matched"


def result(code: str, kind: str, dimension: dict, numerator: int,
           denominator: int | None = None):
    return {"metric_code": code, "dimension_type": kind, "dimension": dimension,
            "numerator": numerator, "denominator": denominator,
            "metric_value": numerator / denominator if denominator else
            (None if denominator == 0 else numerator)}


def calculate(database, month: str, city: str, county_map: Path,
              line_reasons: Path, broadband_file: Path) -> dict:
    city_short = normalize_city(city)
    city_label = display_city(city)
    mapping = load_county_mapping(county_map)
    if city_short not in mapping["standards"]:
        raise ValueError(f"映射表没有地市：{city_label}")
    standards = mapping["standards"][city_short]
    aliases = mapping["aliases"][city_short]
    start, end = integration_opening_period(month)
    source_rows, source_runs = load_rows(database)
    line = calculate_line(source_rows, start, end)
    formal = [row for row in line["details"]["denominator"]
              if normalize_city(text(row.get("地市"))) == city_short]
    line_count, line_rejected, completion, current, previous, unknown = (Counter() for _ in range(6))
    month_start = datetime.strptime(month, "%Y-%m")
    for row in formal:
        county = standardize_county(row.get("区县"), row.get("工单主题"), aliases)
        line_count[county] += 1
        status = text(row.get("工单状态"))
        if status in {"已撤单", "已驳回"}:
            line_rejected[county] += 1
        if status == "已完成":
            completion[county] += 1
            assigned = parse_time(row.get("派单时间"))
            finished = parse_time(row.get("工单结束时间"))
            if assigned is None or assigned > finished:
                unknown[county] += 1
            elif assigned < month_start:
                previous[county] += 1
            elif assigned.year == month_start.year and assigned.month == month_start.month:
                current[county] += 1
            else:
                unknown[county] += 1

    reason_counts = Counter()
    reason_rows = 0
    for _, row in read_sheet(line_reasons, "撤退单清单", {"地市", "区县", "工单主题", "大类"}):
        if normalize_city(text(row["地市"])) != city_short:
            continue
        county = standardize_county(row["区县"], row["工单主题"], aliases)
        reason = LINE_REASONS.get("".join(text(row["大类"]).split()))
        if reason is None:
            raise ValueError(f"专线原因类型未知：{row['大类']}")
        reason_counts[county, reason] += 1
        reason_rows += 1

    accepted, returned, cancelled, broadband_reasons, return_reasons = (Counter() for _ in range(5))
    cell_quality = Counter()
    conflicts = []
    for line_number, row in read_sheet(broadband_file, "Sheet1",
                                       {"CELL", "OP_REMARK", "FIRST_LEVEL_PROBLEM"}):
        source_city, county, state = locate_cell(row["CELL"], mapping["aliases"])
        cell_quality[state] += 1
        if state == "conflict" and len(conflicts) < 30:
            conflicts.append(line_number)
        if source_city != city_short:
            continue
        if row["OP_REMARK"] is None:
            raise ValueError(f"企宽第{line_number}行 OP_REMARK 为空")
        accepted[county] += 1
        remark = "".join(text(row["OP_REMARK"]).split())
        if remark in {"退单完成", "撤单完成"}:
            reason = "".join(text(row["FIRST_LEVEL_PROBLEM"]).split())
            if reason not in BROADBAND_REASONS:
                raise ValueError(f"企宽第{line_number}行原因未知：{reason}")
            broadband_reasons[county, BROADBAND_REASONS[reason]] += 1
            if remark == "退单完成":
                return_reasons[county, BROADBAND_REASONS[reason]] += 1
            (returned if remark == "退单完成" else cancelled)[county] += 1

    codes = {"line_rate": "dedicated_line_opening_withdrawal_rate",
             "line_completed": "dedicated_line_opening_completed_count",
             "line_current": "dedicated_line_opening_completed_current_accepted_count",
             "line_previous": "dedicated_line_opening_completed_previous_accepted_count",
             "line_unknown": "dedicated_line_opening_completed_unknown_accepted_count",
             "qikuan_rate": "qikuan_withdrawal_rate",
             "qikuan_return_rate": "qikuan_return_rate"}
    counties = standards + ([UNASSIGNED_COUNTY] if line_count[UNASSIGNED_COUNTY]
                            or any(c == UNASSIGNED_COUNTY for c, _ in reason_counts) else [])
    rows = []
    for kind, county in [("city", None)] + [("city_county", value) for value in counties]:
        dimension = {"month": month, "city": city_label}
        if county is not None:
            dimension["county"] = county
        value = lambda counter: sum(counter.values()) if county is None else counter[county]
        rows.append(result(codes["line_rate"], kind, dimension, value(line_rejected), value(line_count)))
        for key, counter in [("line_completed", completion), ("line_current", current),
                             ("line_previous", previous), ("line_unknown", unknown)]:
            rows.append(result(codes[key], kind, dimension, value(counter)))
        rows.append(result(codes["qikuan_rate"], kind, dimension,
                           value(returned) + value(cancelled), value(accepted)))
        rows.append(result(codes["qikuan_return_rate"], kind, dimension,
                           value(returned), value(accepted)))
        for name in LINE_REASONS.values():
            rows.append(result(f"line_withdrawal_reason_{name}", kind, dimension,
                               sum(v for (c, r), v in reason_counts.items() if r == name and (county is None or c == county))))
        for name in BROADBAND_REASONS.values():
            rows.append(result(f"qikuan_withdrawal_reason_{name}", kind, dimension,
                               sum(v for (c, r), v in broadband_reasons.items() if r == name and (county is None or c == county))))
            rows.append(result(f"qikuan_return_reason_{name}", kind, dimension,
                               sum(v for (c, r), v in return_reasons.items() if r == name and (county is None or c == county))))
        rows.append(result("qikuan_return_count", kind, dimension, value(returned)))
        rows.append(result("qikuan_cancellation_count", kind, dimension, value(cancelled)))
    if sum(line_count.values()) != sum(line_count[c] for c in counties):
        raise ValueError("专线区县合计与地市不一致")
    if sum(accepted.values()) != sum(accepted[c] for c in standards):
        raise ValueError("企宽区县合计与地市不一致")
    return {"metric_code": OWNER, "metric_version": VERSION, "month": month,
            "city": city_label, "period_start": start, "period_end": end,
            "source_runs": source_runs, "source_files": {
                "line_reasons": str(line_reasons.resolve()),
                "broadband": str(broadband_file.resolve()),
                "county_map": mapping["path"]},
            "source_hashes": {"line_reasons": hashlib.sha256(line_reasons.read_bytes()).hexdigest(),
                              "broadband": hashlib.sha256(broadband_file.read_bytes()).hexdigest(),
                              "county_map": mapping["sha256"]},
            "quality": {"line_rows": len(formal), "line_unassigned": line_count[UNASSIGNED_COUNTY],
                        "line_reason_rows": reason_rows, "cell_match": dict(cell_quality),
                        "cell_conflict_sample_rows": conflicts,
                        "qikuan_city_rows": sum(accepted.values())},
            "results": rows}


def save(database, report: dict) -> str:
    run_id = f"metric_city_support_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}"
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    sources = report["source_runs"] + [f"file:{name}:{hash_value}"
                                      for name, hash_value in report["source_hashes"].items()]
    with connect(database) as connection:
        connection.execute("""INSERT INTO metric_run(metric_run_id,metric_code,metric_version,
            period_start,period_end,started_at,status,source_runs)
            VALUES(?,?,?,?,?,?,'running',?)""",
            (run_id, OWNER, VERSION, report["period_start"], report["period_end"],
             now, json.dumps(sources, ensure_ascii=False)))
    try:
        with connect(database) as connection:
            insert_results(connection, OWNER, run_id, report["results"])
            connection.execute("UPDATE metric_run SET finished_at=?,status='success' WHERE metric_run_id=?",
                               (datetime.now(timezone.utc).isoformat(timespec="seconds"), run_id))
    except Exception as exc:
        with connect(database) as connection:
            connection.execute("UPDATE metric_run SET finished_at=?,status='failed',error_message=? WHERE metric_run_id=?",
                               (datetime.now(timezone.utc).isoformat(timespec="seconds"), str(exc), run_id))
        raise
    return run_id


def run(database, month, city, county_map=DEFAULT_COUNTY_MAP,
        line_reasons=DEFAULT_LINE_REASONS, broadband=DEFAULT_BROADBAND,
        mode="both", output=None):
    initialize(database)
    report = calculate(database, month, city, county_map, line_reasons, broadband)
    report["metric_run_id"] = save(database, report) if mode in {"database", "both"} else None
    if mode in {"file", "both"}:
        path = (output or PROJECT / "outputs" / f"集团撤退单地市指标_{normalize_city(city)}_{month}.json").resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        report["output_file"] = str(path)
    print(f"完成：{report['city']} {month} {report['metric_run_id'] or '仅文件'}", flush=True)
    return report


def main():
    parser = argparse.ArgumentParser(description="计算指定地市集团撤退单及区县指标")
    parser.add_argument("--month", required=True)
    parser.add_argument("--city", required=True)
    parser.add_argument("--database", type=Path)
    parser.add_argument("--county-map", type=Path, default=DEFAULT_COUNTY_MAP)
    parser.add_argument("--line-reasons", type=Path, default=DEFAULT_LINE_REASONS)
    parser.add_argument("--broadband", type=Path, default=DEFAULT_BROADBAND)
    parser.add_argument("--mode", choices=("file", "database", "both"), default="both")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    run(args.database, args.month, args.city, args.county_map, args.line_reasons,
        args.broadband, args.mode, args.output)


if __name__ == "__main__":
    main()
