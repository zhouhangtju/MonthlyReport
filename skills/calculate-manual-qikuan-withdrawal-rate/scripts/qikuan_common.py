from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import load_workbook

CITIES = ["杭州", "嘉兴", "宁波", "温州", "金华", "绍兴", "湖州", "台州", "衢州", "丽水", "舟山"]
CATEGORY_FIELDS = {"用户原因": "customer", "前台原因": "frontDesk", "建设原因": "construction", "其他原因": "other", "网络原因": "network"}


def normalize_text(value: object) -> str:
    return re.sub(r"\s+", "", str(value))


def normalize_city(value: object) -> str:
    city = normalize_text(value)
    return city[:-1] if city.endswith("市") else city


def iter_records(input_path: Path, sheet: str, required_columns: Iterable[str]) -> Iterator[tuple[int, dict]]:
    if not input_path.is_file():
        raise FileNotFoundError(f"找不到输入文件：{input_path}")
    if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("当前只支持 .xlsx 和 .xlsm 文件")
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"找不到工作表 {sheet!r}；可用工作表：{workbook.sheetnames}")
        rows = workbook[sheet].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError(f"工作表 {sheet!r} 为空")
        columns = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
        missing = sorted(set(required_columns) - set(columns))
        if missing:
            raise ValueError(f"工作表 {sheet!r} 缺少必需列：{missing}")
        selected = {name: columns[name] for name in required_columns}
        for row_number, row in enumerate(rows, 2):
            if any(value is not None for value in row):
                yield row_number, {name: row[index] if index < len(row) else None for name, index in selected.items()}
    finally:
        workbook.close()


def validate_month(month: str | None) -> None:
    if month is not None and not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise ValueError("月份格式应为 YYYY-MM")


def reason_summary(counts: Counter, total: int) -> dict:
    return {field: {"count": counts[category], "share": counts[category] / total if total else 0} for category, field in CATEGORY_FIELDS.items()}


def top_three(counters: Counter, denominators: Counter) -> list[dict]:
    rows = [{"city": city, "count": counters[city], "denominator": denominators[city], "rate": counters[city] / denominators[city] if denominators[city] else 0} for city in CITIES]
    return sorted(rows, key=lambda row: (-row["rate"], CITIES.index(row["city"])))[:3]
