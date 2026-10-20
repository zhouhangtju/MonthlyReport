"""手工 Excel 的企宽月末存量；不联网、不修改原始工作簿。"""
from __future__ import annotations

import argparse
import calendar
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

CITIES = "杭州 湖州 嘉兴 金华 丽水 宁波 衢州 绍兴 台州 温州 舟山".split()
VALUE = "日有线宽带出帐用户数"


def periods(month):
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise ValueError("月份必须为 YYYY-MM")
    year, number = map(int, month.split("-"))
    date(year, number, 1)
    date(year - 1, number, 1)
    return {
        "current": month,
        "previous": f"{year if number > 1 else year - 1:04d}-{number - 1 if number > 1 else 12:02d}",
        "year_ago": f"{year - 1:04d}-{number:02d}",
    }


def text(value):
    return "" if value is None else str(value).strip()


def day(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y%m%d")
    value = text(value)
    for fmt in ("%Y%m%d", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, fmt).strftime("%Y%m%d")
        except ValueError:
            pass
    raise ValueError(f"无效日期：{value!r}")


def count(value, location):
    try:
        if isinstance(value, bool):
            raise ValueError()
        number = Decimal(text(value).replace(",", ""))
        if not number.is_finite() or number < 0 or number != number.to_integral_value():
            raise ValueError()
        return int(number)
    except (InvalidOperation, ValueError):
        raise ValueError(f"{location}：出帐用户数必须为非负整数，实际 {value!r}") from None


def read_snapshot(path, month, sheet=None):
    sheet = sheet or month.replace("-", "")
    year, number = map(int, month.split("-"))
    target = date(year, number, calendar.monthrange(year, number)[1]).strftime("%Y%m%d")
    totals, seen, dimensions = Counter(), set(), set()
    selected = raw = 0
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise ValueError(f"{path} 缺少原始工作表 {sheet}")
        rows = book[sheet].iter_rows(values_only=True)
        headers = [text(v) for v in next(rows, ())]
        required = ["日期", "地市", "产品分类", VALUE]
        for name in required:
            if headers.count(name) != 1:
                raise ValueError(f"{sheet} 必须有唯一字段：{name}")
        indices = {name: headers.index(name) for name in required}
        # 前8列为此黄金眼导出的粒度维度，包含两个同名“产品带宽”列。
        expected = ["日期", "地市", "县市", "产品带宽", "产品带宽", "接入技术", "产品分类", "校园属性"]
        if headers[:8] != expected:
            raise ValueError("原始维度布局已变化，需检查后适配脚本")
        for row_number, row in enumerate(rows, 2):
            if all(v is None or text(v) == "" for v in row):
                continue
            raw += 1
            if day(row[indices["日期"]]) != target:
                continue
            product = text(row[indices["产品分类"]])
            if not product:
                raise ValueError(f"{sheet}!第{row_number}行产品分类为空")
            if product != "小微宽带":
                continue
            city = text(row[indices["地市"]]).removesuffix("市")
            if city not in CITIES:
                raise ValueError(f"{sheet}!第{row_number}行未知地市：{city}")
            key = (target, city, *(text(v) for v in row[2:8]))
            if key in dimensions:
                raise ValueError(f"{sheet}!第{row_number}行原始维度重复，需核查是否重复导入")
            dimensions.add(key)
            totals[city] += count(row[indices[VALUE]], f"{sheet}!第{row_number}行")
            seen.add(city)
            selected += 1
        if seen != set(CITIES):
            raise ValueError(f"{target} 小微宽带快照缺少地市：{sorted(set(CITIES) - seen)}")
    finally:
        book.close()
    return {"counts": {city: totals[city] for city in CITIES},
            "audit": {"source_file": str(path.resolve()), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                      "sheet": sheet, "snapshot_date": target, "source_rows": raw,
                      "selected_rows": selected, "product": "小微宽带", "value_field": VALUE}}


def growth(current, baseline):
    if baseline is None:
        return None, "缺少历史文件"
    if baseline == 0:
        return None, "历史基数为0"
    return (current - baseline) / baseline, None


def calculate(month, snapshots):
    expected = periods(month)
    def record(city):
        counts = {}
        for role, snapshot in snapshots.items():
            counts[role] = (sum(snapshot["counts"].values()) if city == "全省"
                            else snapshot["counts"][city]) if snapshot else None
        yoy, yoy_reason = growth(counts["current"], counts["year_ago"])
        mom, mom_reason = growth(counts["current"], counts["previous"])
        return {"city": city, **counts, "yoy": yoy, "mom": mom,
                "yoy_unavailable_reason": yoy_reason, "mom_unavailable_reason": mom_reason}
    province = record("全省")
    cities = [record(city) for city in CITIES]
    return {"schema_version": "1.0", "metric_set": "qikuan_stock", "report_month": month,
            "periods": expected, "province": province, "cities": cities,
            "missing_periods": [expected[k] for k, v in snapshots.items() if v is None],
            "charts": {"province": {"months": [expected[k] for k in ("year_ago", "current", "previous")],
                                   "counts": [province[k] for k in ("year_ago", "current", "previous")]},
                       "cities": {"names": CITIES, "yoy": [v["yoy"] for v in cities],
                                  "mom": [v["mom"] for v in cities]}},
            "audit": {"sources": {k: v["audit"] if v else None for k, v in snapshots.items()},
                      "formula": "(current - baseline) / baseline; province uses summed counts",
                      "generated_at": datetime.now(timezone.utc).isoformat()}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--month", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--previous-input", type=Path)
    parser.add_argument("--previous-sheet")
    parser.add_argument("--year-ago-input", type=Path)
    parser.add_argument("--year-ago-sheet")
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    try:
        months = periods(args.month)
        snapshots = {}
        for role, path, sheet in (("current", args.input, args.sheet),
                                  ("previous", args.previous_input, args.previous_sheet),
                                  ("year_ago", args.year_ago_input, args.year_ago_sheet)):
            if sheet and not path:
                raise ValueError(f"{role} 指定工作表时必须同时提供文件")
            if path and args.output.resolve() == path.resolve():
                raise ValueError("结果路径不能覆盖源文件")
            snapshots[role] = read_snapshot(path, months[role], sheet) if path else None
        result = calculate(args.month, snapshots)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        print(f"全省企宽存量：{result['province']['current']:,}；结果：{args.output.resolve()}")
        if result["missing_periods"]:
            print("缺少比较月份：" + "、".join(result["missing_periods"]))
    except (ValueError, OSError) as exc:
        parser.exit(2, f"错误：{exc}\n")


if __name__ == "__main__":
    main()
