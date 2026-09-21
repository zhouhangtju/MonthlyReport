#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_FIELDS = {"投诉工单流水号", "受理时间", "受理号码", "服务请求类型", "业务地市"}


def text(value) -> str:
    return "" if value is None else str(value).strip()


def parse_time(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    raw = text(value)
    if not raw:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return None


def month_bounds(month: str) -> tuple[datetime, datetime]:
    start = datetime.strptime(month, "%Y-%m")
    end = datetime(start.year + (start.month == 12), 1 if start.month == 12 else start.month + 1, 1)
    return start, end


def normalize_service(value) -> str:
    raw = text(value).replace("->", "→")
    if "政企市场" in raw:
        raw = raw.split("政企市场", 1)[1]
    parts = [part.strip() for part in re.split(r"\s*→\s*", raw) if part.strip()]
    return "→".join(parts)


def read_sheet1(path: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "sheet1" not in workbook.sheetnames:
        raise ValueError(f"{path.name} 缺少 sheet1")
    sheet = workbook["sheet1"]
    header_row = None
    headers = None
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = [text(value) for value in row]
        if "受理时间" in values and "服务请求类型" in values:
            header_row = row_number
            headers = values
            break
    if header_row is None or headers is None:
        raise ValueError(f"{path.name} 未找到表头")
    missing = REQUIRED_FIELDS - set(headers)
    if missing:
        raise ValueError(f"{path.name} 缺少字段: {'、'.join(sorted(missing))}")

    rows = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        if not any(text(value) for value in values):
            continue
        row = {headers[index]: values[index] for index in range(min(len(headers), len(values))) if headers[index]}
        row["_source_file"] = path.name
        row["_source_row"] = source_row
        rows.append(row)
    return rows, {"file": path.name, "sheet": "sheet1", "header_row": header_row, "rows": len(rows)}


def load_inputs(paths: list[Path]) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    rows = []
    sources = []
    for path in paths:
        loaded, source = read_sheet1(path)
        rows.extend(loaded)
        sources.append(source)
    return rows, sources


def audit_row(row: dict[str, object], key: str = "") -> dict[str, object]:
    accepted = row["_accepted_time"]
    return {
        "投诉工单流水号": text(row.get("投诉工单流水号")),
        "受理时间": accepted.strftime("%Y-%m-%d %H:%M:%S") if accepted else "",
        "受理号码": text(row.get("受理号码")),
        "业务地市": text(row.get("业务地市")) or "（空）",
        "重复Key": key,
        "来源文件": row.get("_source_file", ""),
        "来源行": row.get("_source_row", ""),
    }


def rate(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def calculate(source_rows: list[dict[str, object]], month: str) -> dict[str, object]:
    start, end = month_bounds(month)
    valid_enterprise = []
    invalid_time = []
    for source in source_rows:
        if "企业宽带" not in text(source.get("服务请求类型")):
            continue
        row = dict(source)
        row["_accepted_time"] = parse_time(row.get("受理时间"))
        if row["_accepted_time"] is None:
            invalid_time.append({
                "投诉工单流水号": text(row.get("投诉工单流水号")),
                "受理时间原值": text(row.get("受理时间")),
                "来源文件": row.get("_source_file", ""),
                "来源行": row.get("_source_row", ""),
            })
            continue
        row["_service_key"] = normalize_service(row.get("服务请求类型"))
        row["_repeat_key"] = f"{text(row.get('受理号码'))}|{row['_service_key']}"
        row["_city"] = text(row.get("业务地市")) or "（空）"
        valid_enterprise.append(row)

    denominator_rows = [row for row in valid_enterprise if start <= row["_accepted_time"] < end]
    missing_key_rows = [
        row for row in valid_enterprise
        if not text(row.get("受理号码")) or not row["_service_key"]
    ]
    key_rows = [row for row in valid_enterprise if row not in missing_key_rows and row["_accepted_time"] < end]

    by_key = defaultdict(list)
    for row in key_rows:
        by_key[row["_repeat_key"]].append(row)

    removed_within_hour = []
    repeat_rows = []
    one_hour = timedelta(hours=1)
    for key, rows in by_key.items():
        historical = [row for row in rows if row["_accepted_time"] < start]
        target = sorted(
            [row for row in rows if start <= row["_accepted_time"] < end],
            key=lambda row: row["_accepted_time"],
        )
        kept = []
        previous_time = None
        for row in target:
            if previous_time is not None and row["_accepted_time"] - previous_time <= one_hour:
                removed_within_hour.append(audit_row(row, key))
            else:
                kept.append(row)
            previous_time = row["_accepted_time"]
        repeat_rows.extend(kept if historical else kept[1:])

    denominator_by_city = Counter(row["_city"] for row in denominator_rows)
    numerator_by_city = Counter(row["_city"] for row in repeat_rows)
    cities = {
        city: {
            "numerator": numerator_by_city[city],
            "denominator": denominator_by_city[city],
            "rate": rate(numerator_by_city[city], denominator_by_city[city]),
        }
        for city in sorted(denominator_by_city)
    }
    numerator = len(repeat_rows)
    denominator = len(denominator_rows)
    return {
        "metric": "企宽重复投诉率",
        "month": month,
        "rules": {
            "denominator": "目标月受理时间内服务请求类型包含企业宽带的投诉记录数",
            "repeat_key": "受理号码+服务请求类型中政企市场之后的完整路径",
            "within_one_hour": "仅目标月同Key相邻受理时间不超过1小时的后一单从分子候选剔除",
            "numerator": "目标月重复投诉记录数",
            "city": "目标月记录自身的业务地市",
        },
        "province": {"numerator": numerator, "denominator": denominator, "rate": rate(numerator, denominator)},
        "cities": cities,
        "quality": {
            "input_rows": len(source_rows),
            "enterprise_rows_with_valid_time": len(valid_enterprise),
            "target_month_denominator_rows": denominator,
            "within_one_hour_removed_rows": len(removed_within_hour),
            "repeat_rows": numerator,
            "invalid_time_rows": len(invalid_time),
            "missing_repeat_key_rows": len(missing_key_rows),
            "blank_city_denominator_rows": denominator_by_city.get("（空）", 0),
        },
        "audit": {
            "repeat_rows": [audit_row(row, row["_repeat_key"]) for row in repeat_rows],
            "within_one_hour_removed_rows": removed_within_hour,
            "invalid_time_rows": invalid_time,
            "missing_repeat_key_rows": [audit_row(row, row.get("_repeat_key", "")) for row in missing_key_rows],
        },
    }


def write_csv_summary(result: dict[str, object], path: Path) -> None:
    fields = ["范围", "重复投诉数", "企业宽带投诉数", "重复投诉率"]
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        province = result["province"]
        writer.writerow({
            "范围": "全省", "重复投诉数": province["numerator"],
            "企业宽带投诉数": province["denominator"],
            "重复投诉率": "" if province["rate"] is None else f"{province['rate']:.2%}",
        })
        for city, values in result["cities"].items():
            writer.writerow({
                "范围": city, "重复投诉数": values["numerator"],
                "企业宽带投诉数": values["denominator"],
                "重复投诉率": "" if values["rate"] is None else f"{values['rate']:.2%}",
            })


def main() -> None:
    parser = argparse.ArgumentParser(description="合并月度大音投诉Excel并计算企宽重复投诉率")
    parser.add_argument("--input", action="append", required=True, type=Path, help="可重复传入多个连续月份Excel")
    parser.add_argument("--month", required=True, help="目标月份，格式YYYY-MM")
    parser.add_argument("--json-output", required=True, type=Path)
    parser.add_argument("--csv-output", type=Path)
    args = parser.parse_args()

    rows, sources = load_inputs(args.input)
    result = calculate(rows, args.month)
    result["sources"] = sources
    result.update(schema_version='1.0', metric_set='qikuan_repeat_complaint_rate', report_month=args.month)
    args.json_output.parent.mkdir(parents=True, exist_ok=True)
    args.json_output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.csv_output:
        args.csv_output.parent.mkdir(parents=True, exist_ok=True)
        write_csv_summary(result, args.csv_output)


if __name__ == "__main__":
    main()
