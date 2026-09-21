"""Inspect source headers without copying customer records into the repository."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from openpyxl import load_workbook

PATTERNS = {
    "terminal_material_names": "物料名称表",
    "integration_removal_order": "一体化专线拆机清单_",
    "eoms_service_removal_order": "服务类产品支撑工单_",
    "integration_terminal_outbound": "终端出库_",
    "integration_terminal_inbound": "终端入库_",
    "integration_material_baseline": "全省物资基准库",
    "orch_opening": "专线工单_",
    "orch_install": "互联网专线新装单_",
    "integration_opening": "售中开通工单_",
    "eoms_complaint": "投诉工单_原始数据_",
    "youshu_install": "企宽新装清单_",
    "youshu_complaint": "企宽投诉清单_batch",
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    files = [p for p in args.directory.rglob("*.xlsx")
             if not p.name.startswith(("~$", "._")) and "__MACOSX" not in p.parts and "合并" not in p.name]
    manifest = {}
    for code, prefix in PATTERNS.items():
        choices = sorted((p for p in files if p.name.startswith(prefix)), key=lambda p: (-p.stat().st_size, len(p.parts), p.name))
        if not choices:
            raise ValueError(f"Missing source file: {code}")
        source = choices[0]
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            workbook.worksheets[0].reset_dimensions()
            rows = workbook.worksheets[0].iter_rows(values_only=True)
            headers = list(next(rows))
            samples = [next(rows, None) for _ in range(100)]
            manifest[code] = {"source_file": source.name, "headers": headers}
            print(json.dumps({"dataset": code, "headers": headers,
                              "sample_types": sorted({type(v).__name__ for row in samples if row for v in row if v is not None})}, ensure_ascii=True), flush=True)
        finally:
            workbook.close()
    args.output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()

