#!/usr/bin/env python3
"""Inspect candidate key completeness and uniqueness in exported Excel files."""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def inspect(path: Path, columns: list[str]) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("工作表为空")

        headers = [str(value).strip() if value is not None else "" for value in header_row]
        indexes = {column: headers.index(column) for column in columns if column in headers}
        missing = [column for column in columns if column not in indexes]
        values: dict[str, set[str]] = defaultdict(set)
        nonempty: dict[str, int] = defaultdict(int)
        row_count = 0

        for row in rows:
            row_count += 1
            for column, index in indexes.items():
                value = "" if index >= len(row) or row[index] is None else str(row[index]).strip()
                if value:
                    nonempty[column] += 1
                    values[column].add(value)

        print(f"文件：{path}")
        print(f"数据行数：{row_count}")
        if missing:
            print(f"缺少候选字段：{', '.join(missing)}")
        for column in indexes:
            unique_count = len(values[column])
            print(
                f"{column}: 非空={nonempty[column]}, 唯一值={unique_count}, "
                f"重复行={nonempty[column] - unique_count}, 空值={row_count - nonempty[column]}"
            )
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="检查 Excel 候选主键的完整性和唯一性")
    parser.add_argument("path", type=Path)
    parser.add_argument("columns", nargs="+")
    args = parser.parse_args()
    inspect(args.path, args.columns)


if __name__ == "__main__":
    main()
